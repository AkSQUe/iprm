"""XLSX: шаблон витратних матеріалів MM Medic і звіт по резервах.

Єдиний домен тут без імпорту в БД: шаблон віддається партнеру порожнім,
а назад приймається лише пара «артикул -> кількість».
"""

from __future__ import annotations

import io
import time

from datetime import date
from pathlib import Path

from openpyxl import (
    Workbook,
    load_workbook,
)
from openpyxl.utils import get_column_letter

from ._common import (
    _apply_table_style,
    _apply_zebra,
    _find_sheet,
    _int,
    _read_sheet,
    _set_column_widths,
    _style_header,
    logger,
    write_cell,
)


# ==================== MM MEDIC MATERIALS TEMPLATE ====================

_MATERIALS_COLS = ['image', 'sku', 'name', 'available', 'quantity']
_MATERIALS_LABELS = {
    'image': 'Зображення',
    'sku': 'Артикул',
    'name': 'Назва',
    'available': 'Наявно',
    'quantity': 'Кількість',
}
_MATERIALS_WIDTHS = {'image': 9, 'sku': 20, 'name': 46, 'available': 12, 'quantity': 14}
_MATERIALS_THUMB_PX = 40


# Сумарний бюджет на ВСІ мініатюри одного експорту. Без нього каталог на
# сотню позицій міг тягнутись сотні секунд (до 8 с на позицію) і впертись у
# таймаут шлюзу -- заради декоративних картинок.
_MATERIALS_THUMB_BUDGET_SECONDS = 20.0
_MATERIALS_THUMB_MAX_BYTES = 2_000_000


def _download_thumb(url, max_px=_MATERIALS_THUMB_PX):
    """Завантажити зображення товару й повернути BytesIO з PNG-мініатюрою
    (max_px), або None (порожнє/не-http/збій/не зображення). Best-effort."""
    url = (url or '').strip()
    if not url.startswith(('http://', 'https://')):
        return None
    try:
        import requests
        from PIL import Image as PILImage

        # stream + порізний ліміт: раніше resp.content матеріалізував тіло
        # ЦІЛКОМ, і перевірка розміру після цього вже нічого не рятувала.
        with requests.get(url, timeout=(3.0, 5.0), stream=True) as resp:
            if not resp.ok:
                return None
            chunks, total = [], 0
            for chunk in resp.iter_content(64 * 1024):
                total += len(chunk)
                if total > _MATERIALS_THUMB_MAX_BYTES:
                    return None
                chunks.append(chunk)
        img = PILImage.open(io.BytesIO(b''.join(chunks)))
        img.thumbnail((max_px, max_px))
        out = io.BytesIO()
        img.convert('RGB').save(out, format='PNG')
        out.seek(0)
        return out
    except Exception:
        logger.info('materials thumb fetch failed: %s', url)
        return None


def export_materials_template_xlsx(catalog: list[dict]) -> io.BytesIO:
    """Шаблон для резервування витратних матеріалів MM Medic.

    `catalog` -- список dict з ключами sku, name, available, image (як віддає
    MM Medic /catalog). Перша колонка `Зображення` містить вбудовану мініатюру
    товару (не URL); зображення тягнуться з MM Medic best-effort. Колонка
    `Кількість` порожня: адмін вписує потрібні кількості; незаповнені рядки на
    імпорті ігноруються.
    """
    from openpyxl.drawing.image import Image as XLImage

    wb = Workbook()
    ws = wb.active
    ws.title = 'Матеріали'
    _style_header(ws, _MATERIALS_COLS, _MATERIALS_LABELS)

    img_col = get_column_letter(_MATERIALS_COLS.index('image') + 1)
    sku_i = _MATERIALS_COLS.index('sku') + 1
    name_i = _MATERIALS_COLS.index('name') + 1
    avail_i = _MATERIALS_COLS.index('available') + 1

    deadline = time.monotonic() + _MATERIALS_THUMB_BUDGET_SECONDS
    skipped_thumbs = 0

    for row_idx, item in enumerate(catalog or [], start=2):
        ws.cell(row=row_idx, column=sku_i, value=item.get('sku') or '')
        ws.cell(row=row_idx, column=name_i, value=item.get('name') or '')
        ws.cell(row=row_idx, column=avail_i, value=item.get('available'))
        # quantity column left empty for the admin

        # Вичерпали бюджет -- решта рядків без картинок. Файл лишається
        # придатним до роботи, а адмін не чекає таймауту.
        if time.monotonic() >= deadline:
            skipped_thumbs += 1
            continue
        thumb = _download_thumb(item.get('image'))
        if thumb is not None:
            xi = XLImage(thumb)
            xi.width = _MATERIALS_THUMB_PX
            xi.height = _MATERIALS_THUMB_PX
            ws.add_image(xi, f'{img_col}{row_idx}')
            ws.row_dimensions[row_idx].height = 32

    if skipped_thumbs:
        logger.info('materials export: %s thumbnails skipped (time budget)',
                    skipped_thumbs)

    last_row = ws.max_row
    _set_column_widths(ws, _MATERIALS_COLS, _MATERIALS_WIDTHS)
    _apply_zebra(ws, len(_MATERIALS_COLS), first_data_row=2, last_data_row=last_row)
    if last_row >= 2:
        _apply_table_style(ws, _MATERIALS_COLS, 'tblMaterials', last_row)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


_RESV_COLS = ['event', 'date', 'status', 'positions', 'reserved', 'actual']
_RESV_LABELS = {
    'event': 'Захід', 'date': 'Дата', 'status': 'Статус',
    'positions': 'Позицій', 'reserved': 'Зарезервовано', 'actual': 'Фактично',
}
_RESV_WIDTHS = {'event': 46, 'date': 14, 'status': 16, 'positions': 10,
                'reserved': 14, 'actual': 12}


def export_material_reservations_xlsx(reservations) -> io.BytesIO:
    """Огляд резервувань матеріалів -> xlsx (для експорту зі сторінки огляду)."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Резервування'
    _style_header(ws, _RESV_COLS, _RESV_LABELS)

    for row_idx, r in enumerate(reservations, start=2):
        course = r.instance.course.title if (r.instance and r.instance.course) else '—'
        date = (r.instance.start_date.strftime('%d.%m.%Y')
                if (r.instance and r.instance.start_date) else '')
        reserved = sum((it.quantity_reserved or 0) for it in r.items)
        actual = sum((it.quantity_actual or 0) for it in r.items)
        values = [course, date, r.status_label, len(r.items), reserved, actual]
        for col_idx, v in enumerate(values, start=1):
            write_cell(ws, row_idx, col_idx, v)

    last_row = ws.max_row
    _set_column_widths(ws, _RESV_COLS, _RESV_WIDTHS)
    _apply_zebra(ws, len(_RESV_COLS), first_data_row=2, last_data_row=last_row)
    if last_row >= 2:
        _apply_table_style(ws, _RESV_COLS, 'tblResv', last_row)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


def parse_materials_xlsx(path: Path) -> dict[str, int]:
    """Прочитати заповнений шаблон -> {sku: quantity} лише для quantity > 0.

    Рядки з порожньою/нульовою/невалідною кількістю ігноруються (за вимогою).
    """
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        # Раніше тут стояло _find_sheet(wb, 'Матеріали') -- ключа 'Матеріали'
        # у SHEET_ALIASES немає, тож функція ЗАВЖДИ повертала None і код
        # мовчки читав активний лист. Варто було адміну лишити активним
        # інший лист (напр. власні нотатки), як імпорт повертав порожньо, а
        # повідомлення казало "Завантажено 0 позицій".
        ws = _find_sheet(wb, 'materials')
        if ws is None:
            if len(wb.sheetnames) == 1:
                ws = wb[wb.sheetnames[0]]  # файл зібрали вручну -- беремо єдиний
            else:
                raise ValueError(
                    'у файлі немає листа "Матеріали"; наявні листи: '
                    + ', '.join(wb.sheetnames)
                )
        rows = _read_sheet(ws, ['sku', 'quantity'], _MATERIALS_LABELS)
    finally:
        # read_only mode keeps the file handle open; must close or Windows
        # blocks the subsequent cleanup_upload() unlink.
        wb.close()

    result: dict[str, int] = {}
    for row in rows:
        sku = (str(row.get('sku')).strip() if row.get('sku') is not None else '')
        if not sku:
            continue
        try:
            qty = _int(row.get('quantity'))
        except ValueError:
            continue  # невалідне число -> ігнор
        if not qty or qty <= 0:
            continue
        result[sku] = qty
    return result
