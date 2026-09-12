"""XLSX: учасники заходів (реєстрації).

Найчутливіший з аркушів: у ньому персональні дані, і саме сюди потрапляє
чужий текст -- місце роботи, нотатки. Через це запис клітинок іде через
write_cell, який знімає тип формули: інакше учасник міг би керувати
вмістом файлу, який відкриє менеджер.

Запис не робиться прямо в модель -- він проходить через
participant_service.upsert_participant, щоб xlsx-імпорт і ручне
редагування в адмінці не розійшлися в правилах.
"""

from __future__ import annotations

import io
import re

from dataclasses import (
    dataclass,
    field,
)
from datetime import (
    date,
    datetime,
)
from pathlib import Path

from openpyxl import (
    Workbook,
    load_workbook,
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy.orm import joinedload

from app.extensions import db
from app.models.course_instance import CourseInstance
from app.models.medical_profile import MedicalProfile
from app.models.registration import EventRegistration
from app.data.specializations import SPECIALIZATIONS
from app.models.user import User

from ._common import (
    WRAP,
    _APPLY_FAILED_MESSAGE,
    _DROPDOWN_BUFFER_ROWS,
    _add_inline_dropdown,
    _apply_number_formats,
    _apply_table_style,
    _apply_zebra,
    _bool,
    _decimal,
    _fill,
    _find_sheet,
    _from_lines,
    _int,
    _points_cell,
    _read_sheet,
    _set_column_widths,
    _str,
    _style_header,
    logger,
    record_row_error,
    write_cell,
)


# ======================================================================
# PARTICIPANTS (учасники заходів)
# ======================================================================

PARTICIPANT_COLS = [
    'reg_id', 'event', 'last_name', 'first_name', 'middle_name', 'email',
    'phone', 'participant_type', 'birth_date', 'education', 'workplace',
    'position', 'specializations', 'status', 'payment_status',
    'payment_amount', 'attended', 'cpd_points_awarded', 'participation_format',
    'experience_years', 'license_number', 'admin_notes',
]

PARTICIPANT_LABELS = {
    'reg_id': 'ID реєстрації',
    'event': 'Захід',
    'last_name': 'Прізвище',
    'first_name': "Ім'я",
    'middle_name': 'По батькові',
    'email': 'Email',
    'phone': 'Телефон',
    'participant_type': 'Тип учасника',
    'birth_date': 'Дата народження',
    'education': 'Освіта',
    'workplace': 'Місце роботи / місто',
    'position': 'Посада',
    'specializations': 'Спеціалізації',
    'status': 'Статус',
    'payment_status': 'Оплата',
    'payment_amount': 'Сума (грн)',
    'attended': 'Присутній',
    'cpd_points_awarded': 'Бали БПР',
    'participation_format': 'Формат участі',
    'experience_years': 'Стаж (років)',
    'license_number': 'Ліцензія',
    'admin_notes': 'Нотатки',
    'promo_code': 'Промокод',
    'discount_amount': 'Знижка (грн)',
}

# Колонки ЛИШЕ на вивантаження: промокод і знижка -- довідкові, ними
# володіє promo_service, тож імпорт їх не читає (невідомі заголовки
# _read_sheet просто ігнорує, тож round-trip "вивантажив -> завантажив"
# лишається робочим).
PARTICIPANT_EXPORT_COLS = PARTICIPANT_COLS + ['promo_code', 'discount_amount']

PARTICIPANT_WIDTHS = {
    'reg_id': 12,
    'event': 46,
    'last_name': 20,
    'first_name': 18,
    'middle_name': 20,
    'email': 30,
    'phone': 18,
    'participant_type': 28,
    'birth_date': 16,
    'education': 40,
    'workplace': 34,
    'position': 26,
    'specializations': 40,
    'status': 16,
    'payment_status': 16,
    'payment_amount': 14,
    'attended': 12,
    'cpd_points_awarded': 12,
    'participation_format': 16,
    'experience_years': 12,
    'license_number': 18,
    'admin_notes': 40,
    'promo_code': 18,
    'discount_amount': 14,
}

REG_STATUS_LABEL = dict(EventRegistration.STATUSES)
REG_STATUS_KEY_BY_LABEL = {v: k for k, v in REG_STATUS_LABEL.items()}
PAYMENT_STATUS_LABEL = dict(EventRegistration.PAYMENT_STATUSES)
PAYMENT_STATUS_KEY_BY_LABEL = {v: k for k, v in PAYMENT_STATUS_LABEL.items()}
PARTICIPANT_TYPE_LABEL = dict(MedicalProfile.PARTICIPANT_TYPES)
PARTICIPANT_TYPE_KEY_BY_LABEL = {v: k for k, v in PARTICIPANT_TYPE_LABEL.items()}
SPEC_LABEL_BY_CODE = dict(SPECIALIZATIONS)
SPEC_CODE_BY_LABEL = {v: k for k, v in SPECIALIZATIONS}

VALID_REG_STATUSES = set(REG_STATUS_LABEL.keys())
VALID_PAYMENT_STATUSES = set(PAYMENT_STATUS_LABEL.keys())
VALID_PARTICIPANT_TYPES = set(PARTICIPANT_TYPE_LABEL.keys())

# Формат участі -- окремий (звужений) словник, а не FORMAT_KEY_BY_LABEL
# заходу: там є ще й 'Гібрид', який для participation_format недопустимий
# (це поле лише online/offline/NULL -- сам гібрид визначається на рівні
# заходу). Ключі -- у нижньому регістрі, бо порівняння регістронезалежне.
_PARTICIPATION_FORMAT_KEY_BY_TEXT = {
    'online': 'online', 'offline': 'offline',
    'онлайн': 'online', 'офлайн': 'offline',
}


def _parse_participation_format(raw):
    """Формат участі з комірки -> 'online' / 'offline' / None.

    Порожнє значення -- коректний стан (None): далі підхопить
    effective_participation_format (тариф -> формат заходу). Приймає і
    внутрішній код, і українську назву в будь-якому регістрі. Будь-що інше
    -- зрозуміла per-row помилка тут, а не падіння на DB CHECK-констрейнті
    у apply.
    """
    text = _str(raw)
    if not text:
        return None
    key = _PARTICIPATION_FORMAT_KEY_BY_TEXT.get(text.strip().lower())
    if key is None:
        raise ValueError(
            f'формат участі {raw!r} – допустимі: Онлайн, Офлайн (або порожньо)'
        )
    return key


REG_STATUS_FILLS = {
    'pending': _fill('FEF3C7'),     # yellow
    'confirmed': _fill('DBEAFE'),   # blue
    'completed': _fill('A7F3D0'),   # green
    'cancelled': _fill('FECACA'),   # red
}
PAYMENT_STATUS_FILLS = {
    'unpaid': _fill('F3F4F6'),      # gray
    'pending': _fill('FEF3C7'),     # yellow
    'paid': _fill('A7F3D0'),        # green
    'refunded': _fill('FECACA'),    # red
}

_EVENTS_SHEET_NAME = 'Заходи'
_SPEC_SHEET_NAME = 'Спеціалізації (довідник)'


def _date(v) -> date | None:
    """Прийняти date/datetime (openpyxl) або рядок (ISO / dd.mm.yyyy)."""
    if v is None or v == '':
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ('%Y-%m-%d', '%d.%m.%Y', '%d/%m/%Y'):
        try:
            return datetime.strptime(s[:10], fmt).date()
        except ValueError:
            continue
    raise ValueError(f'неможливо розпарсити дату: {v!r}')


def _participant_event_label(instance) -> str:
    """Людиночитний ярлик заходу для колонки/довідника. Починається з
    '#<id>' -> id можна розпарсити навіть якщо назву трохи змінили.

    Єдина реалізація живе у participant_service.event_label (спільна з
    випадними списками форми)."""
    from app.services import participant_service
    return participant_service.event_label(instance, with_id=True)


def _resolve_event_to_id(label, label_to_id, instance_by_id):
    """Ярлик заходу -> instance_id. Спершу точний збіг, потім '#<id>'."""
    if not label:
        return None
    if label in label_to_id:
        return label_to_id[label]
    m = re.match(r'^\s*#(\d+)', str(label))
    if m:
        iid = int(m.group(1))
        if iid in instance_by_id:
            return iid
    return None


def _parse_spec_cell(raw):
    """Текст спеціалізацій (по рядку / через кому) -> (codes, unknown).

    Приймає і label ('Терапія'), і code ('therapy')."""
    codes, unknown, seen = [], [], set()
    for line in _from_lines(raw):
        for part in re.split(r'[;,]', line):
            p = part.strip()
            if not p:
                continue
            if p in SPEC_LABEL_BY_CODE:
                code = p
            elif p in SPEC_CODE_BY_LABEL:
                code = SPEC_CODE_BY_LABEL[p]
            else:
                unknown.append(p)
                continue
            if code not in seen:
                seen.add(code)
                codes.append(code)
    return codes, unknown


def _add_events_sheet(wb) -> int:
    """Reference-sheet із заходами. Ярлик у колонці A (для drop-down),
    id у B. Повертає номер останнього рядка з даними."""
    ws = wb.create_sheet(_EVENTS_SHEET_NAME)
    cols = ['event', 'id']
    _style_header(ws, cols, {'event': 'Захід (значення)', 'id': 'ID'})
    instances = (
        CourseInstance.query
        .options(joinedload(CourseInstance.course))
        .order_by(CourseInstance.start_date.desc().nullslast())
        .all()
    )
    for row_idx, inst in enumerate(instances, start=2):
        ws.cell(row=row_idx, column=1, value=_participant_event_label(inst)).alignment = WRAP
        ws.cell(row=row_idx, column=2, value=inst.id)
    _set_column_widths(ws, cols, {'event': 70, 'id': 8})
    _apply_zebra(ws, len(cols), first_data_row=2, last_data_row=1 + len(instances))
    _apply_table_style(ws, cols, 'tblEvents', last_data_row=1 + len(instances))
    return 1 + len(instances)


def _add_specializations_sheet(wb) -> None:
    """Reference-sheet (label + code) для ручного пошуку спеціалізацій."""
    ws = wb.create_sheet(_SPEC_SHEET_NAME)
    cols = ['label', 'code']
    _style_header(ws, cols, {'label': 'Спеціалізація', 'code': 'Код'})
    for row_idx, (code, label) in enumerate(SPECIALIZATIONS, start=2):
        ws.cell(row=row_idx, column=1, value=label).alignment = WRAP
        ws.cell(row=row_idx, column=2, value=code)
    _set_column_widths(ws, cols, {'label': 50, 'code': 30})
    _apply_zebra(ws, len(cols), first_data_row=2, last_data_row=1 + len(SPECIALIZATIONS))
    _apply_table_style(ws, cols, 'tblSpecs', last_data_row=1 + len(SPECIALIZATIONS))


def _add_ref_dropdown(ws, column_key, columns, last_data_row, sheet_name,
                      ref_last_row, title='', prompt=''):
    """Drop-down з reference-sheet (колонка A) на вказану колонку."""
    if ref_last_row < 2:
        return
    col_letter = get_column_letter(columns.index(column_key) + 1)
    formula = f"='{sheet_name}'!$A$2:$A${ref_last_row}"
    dv = DataValidation(
        type='list', formula1=formula, allow_blank=True,
        showDropDown=False, errorStyle='warning',
        error='Значення відсутнє у довіднику.', errorTitle='Невідоме значення',
        prompt=prompt, promptTitle=title,
    )
    final_row = max(last_data_row, 1) + _DROPDOWN_BUFFER_ROWS
    dv.add(f'{col_letter}2:{col_letter}{final_row}')
    ws.add_data_validation(dv)


@dataclass
class ParticipantChange:
    action: str  # 'create' | 'update' | 'unchanged' | 'error'
    name: str = ''
    event: str = ''
    fields_changed: list[str] = field(default_factory=list)
    error: str | None = None


@dataclass
class ParticipantsImportPlan:
    participants: list[dict] = field(default_factory=list)
    changes: list[ParticipantChange] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    # (ключ аркуша, номер рядка, текст) -- для повернення помилок
    # у копію завантаженого файлу.
    row_errors: list[tuple[str, int, str]] = field(default_factory=list)

    @property
    def is_valid(self) -> bool:
        return not self.errors

    @property
    def counts(self) -> dict[str, int]:
        c = {'create': 0, 'update': 0, 'unchanged': 0, 'error': 0}
        for ch in self.changes:
            c[ch.action] = c.get(ch.action, 0) + 1
        return c


def export_participants_xlsx(instance_id=None, blank=False) -> io.BytesIO:
    """Експорт учасників у xlsx (також слугує формою-шаблоном для додавання).

    instance_id: якщо задано -- лише учасники цього заходу.
    blank: True -- порожній шаблон (лише заголовки + dropdown-и, без даних).
    Reference-sheet «Заходи» + drop-down дозволяють додавати рядки для
    будь-якого заходу.

    Шаблон для заповнення (blank) містить лише редаговані колонки, а
    вивантаження даних -- ще й довідкові промокод/знижку: у шаблоні вони
    були б пасткою (заповнив -- нічого не сталось), бо знижками володіє
    promo_service, і імпорт їх не читає.
    """
    cols = PARTICIPANT_COLS if blank else PARTICIPANT_EXPORT_COLS
    wb = Workbook()
    ws = wb.active
    ws.title = 'Учасники'
    _style_header(ws, cols, PARTICIPANT_LABELS)

    from app.services import participant_service

    if blank:
        regs = []
    else:
        q = (
            EventRegistration.query
            .options(
                joinedload(EventRegistration.user).joinedload(User.medical_profile),
                joinedload(EventRegistration.instance).joinedload(CourseInstance.course),
            )
            .order_by(EventRegistration.created_at.desc())
        )
        if instance_id:
            q = q.filter(EventRegistration.instance_id == instance_id)
        regs = q.all()

    for row_idx, reg in enumerate(regs, start=2):
        user = reg.user
        profile = user.medical_profile if user else None
        raw_email = user.email if user else ''
        email = '' if participant_service.is_placeholder_email(raw_email) else (raw_email or '')
        spec_codes = (profile.specializations if profile else []) or []
        spec_text = '\n'.join(SPEC_LABEL_BY_CODE.get(c, c) for c in spec_codes)
        ptype = profile.participant_type if profile else None
        values = [
            reg.id,
            _participant_event_label(reg.instance) if reg.instance else '',
            (user.last_name if user else '') or '',
            (user.first_name if user else '') or '',
            (profile.middle_name if profile else '') or '',
            email,
            reg.phone or (profile.phone if profile else '') or '',
            PARTICIPANT_TYPE_LABEL.get(ptype, '') if ptype else '',
            profile.birth_date if profile else None,
            (profile.education if profile else '') or '',
            (profile.workplace if profile else '') or reg.workplace or '',
            (profile.position if profile else '') or '',
            spec_text,
            REG_STATUS_LABEL.get(reg.status, reg.status or ''),
            PAYMENT_STATUS_LABEL.get(reg.payment_status, reg.payment_status or ''),
            float(reg.payment_amount) if reg.payment_amount is not None else None,
            'Так' if reg.attended else 'Ні',
            reg.cpd_points_awarded,
            # Сира колонка, а не effective_participation_format: імпорт кладе
            # значення саме туди, тож віддавати обчислене з тарифу означало б
            # круговоротом "вивантажив -> виправив прізвище -> завантажив"
            # жорстко фіксувати похідне значення в кожному рядку (подальша
            # зміна тарифу вже не рухала б бали) і позначати "формат участі"
            # зміненим у кожному рядку, де колонка була NULL. Порожньо тут --
            # "за тарифом", так само як у формі картки учасника.
            {'online': 'Онлайн', 'offline': 'Офлайн'}.get(
                reg.participation_format, ''),
            reg.experience_years,
            reg.license_number or '',
            reg.admin_notes or '',
            reg.promo_code.code if reg.promo_code else '',
            float(reg.discount_amount) if reg.discount_amount is not None else None,
        ]
        for col_idx, v in enumerate(values, start=1):
            cell = write_cell(ws, row_idx, col_idx, v)
            cell.alignment = WRAP

    last_row = ws.max_row
    _apply_zebra(ws, len(cols), first_data_row=2, last_data_row=last_row)

    # Кольори за статусом / оплатою.
    st_col = PARTICIPANT_COLS.index('status') + 1
    pay_col = PARTICIPANT_COLS.index('payment_status') + 1
    for row_idx, reg in enumerate(regs, start=2):
        if reg.status in REG_STATUS_FILLS:
            ws.cell(row=row_idx, column=st_col).fill = REG_STATUS_FILLS[reg.status]
        if reg.payment_status in PAYMENT_STATUS_FILLS:
            ws.cell(row=row_idx, column=pay_col).fill = PAYMENT_STATUS_FILLS[reg.payment_status]

    _set_column_widths(ws, cols, PARTICIPANT_WIDTHS)
    _apply_number_formats(ws, cols, last_row)

    # Reference-sheets + drop-downs.
    events_last_row = _add_events_sheet(wb)
    _add_specializations_sheet(wb)
    _add_ref_dropdown(
        ws, 'event', PARTICIPANT_COLS, last_row,
        _EVENTS_SHEET_NAME, events_last_row,
        title='Захід', prompt='Оберіть захід зі списку (натисніть стрілочку)',
    )
    _add_inline_dropdown(
        ws, 'participant_type', PARTICIPANT_COLS,
        options=[label for _k, label in MedicalProfile.PARTICIPANT_TYPES],
        last_data_row=last_row, title='Тип учасника',
        hint='Лікар / Молодший спеціаліст / Інтерн / Студент',
    )
    _add_inline_dropdown(
        ws, 'status', PARTICIPANT_COLS,
        options=[label for _k, label in EventRegistration.STATUSES],
        last_data_row=last_row, title='Статус',
        hint='Очікує / Підтверджено / Скасовано / Завершено',
    )
    _add_inline_dropdown(
        ws, 'payment_status', PARTICIPANT_COLS,
        options=[label for _k, label in EventRegistration.PAYMENT_STATUSES],
        last_data_row=last_row, title='Оплата',
        hint='Не оплачено / Очікує оплати / Оплачено / Повернено',
    )
    _add_inline_dropdown(
        ws, 'attended', PARTICIPANT_COLS, options=['Так', 'Ні'],
        last_data_row=last_row, title='Присутній', hint='Так / Ні',
    )

    _apply_table_style(ws, cols, 'tblParticipants', last_row)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


_PARTICIPANT_DIFF_LABELS = {
    'last_name': 'прізвище', 'first_name': "ім'я", 'middle_name': 'по батькові',
    'email': 'email', 'phone': 'телефон', 'participant_type': 'тип учасника',
    'birth_date': 'дата народж.', 'education': 'освіта', 'workplace': 'місце роботи',
    'position': 'посада', 'specializations': 'спеціалізації', 'status': 'статус',
    'payment_status': 'оплата', 'payment_amount': 'сума', 'attended': 'присутність',
    'cpd_points_awarded': 'бали БПР', 'participation_format': 'формат участі',
    'experience_years': 'стаж',
    'license_number': 'ліцензія', 'admin_notes': 'нотатки',
}


def _diff_participant(reg, data):
    """Які поля зміняться при застосуванні data до наявної реєстрації.

    Дзеркалить семантику participant_service.upsert_participant: identity та
    профіль оновлюються лише непорожніми значеннями; реєстраційні поля --
    завжди (replace). Повертає список людиночитних назв змінених полів
    (порожній -> рядок 'без змін')."""
    user = reg.user
    profile = user.medical_profile if user else None
    changed = []

    def norm(v):
        return v if v not in ('', None) else None

    # Identity + профіль: оновлюються лише непорожнім вводом.
    soft = [
        ('last_name', user.last_name if user else None),
        ('first_name', user.first_name if user else None),
        ('middle_name', profile.middle_name if profile else None),
        ('participant_type', profile.participant_type if profile else None),
        ('birth_date', profile.birth_date if profile else None),
        ('education', profile.education if profile else None),
        ('workplace', profile.workplace if profile else None),
        ('position', profile.position if profile else None),
    ]
    for key, cur in soft:
        new = data.get(key)
        if new in ('', None):
            continue
        if norm(new) != norm(cur):
            changed.append(_PARTICIPANT_DIFF_LABELS[key])

    new_email = (data.get('email') or '').strip().lower()
    if new_email and user and new_email != (user.email or '').lower():
        changed.append(_PARTICIPANT_DIFF_LABELS['email'])

    new_specs = list(data.get('specializations') or [])
    if new_specs and new_specs != list((profile.specializations if profile else []) or []):
        changed.append(_PARTICIPANT_DIFF_LABELS['specializations'])

    # Реєстраційні поля -- replace-семантика (порівнюємо завжди).
    if norm(data.get('phone')) != norm(reg.phone):
        changed.append(_PARTICIPANT_DIFF_LABELS['phone'])
    if (data.get('status') or 'confirmed') != reg.status:
        changed.append(_PARTICIPANT_DIFF_LABELS['status'])
    if (data.get('payment_status') or 'unpaid') != reg.payment_status:
        changed.append(_PARTICIPANT_DIFF_LABELS['payment_status'])
    if bool(data.get('attended')) != bool(reg.attended):
        changed.append(_PARTICIPANT_DIFF_LABELS['attended'])
    if data.get('cpd_points_awarded') != reg.cpd_points_awarded:
        changed.append(_PARTICIPANT_DIFF_LABELS['cpd_points_awarded'])
    if data.get('participation_format') != reg.participation_format:
        changed.append(_PARTICIPANT_DIFF_LABELS['participation_format'])
    if data.get('experience_years') != reg.experience_years:
        changed.append(_PARTICIPANT_DIFF_LABELS['experience_years'])
    if norm(data.get('license_number')) != norm(reg.license_number):
        changed.append(_PARTICIPANT_DIFF_LABELS['license_number'])
    if norm(data.get('admin_notes')) != norm(reg.admin_notes):
        changed.append(_PARTICIPANT_DIFF_LABELS['admin_notes'])
    pa_new, pa_cur = data.get('payment_amount'), reg.payment_amount
    if (pa_new is None) != (pa_cur is None) or (
        pa_new is not None and pa_cur is not None and pa_new != pa_cur
    ):
        changed.append(_PARTICIPANT_DIFF_LABELS['payment_amount'])

    return changed


def parse_participants_xlsx(path: Path) -> ParticipantsImportPlan:
    plan = ParticipantsImportPlan()
    try:
        wb = load_workbook(filename=str(path), read_only=False, data_only=True)
    except Exception as exc:
        plan.errors.append(f'Не вдалося відкрити xlsx: {exc}')
        return plan

    ws = _find_sheet(wb, 'participants')
    if ws is None:
        plan.errors.append('Відсутній sheet "Учасники"')
        return plan

    try:
        rows = _read_sheet(ws, PARTICIPANT_COLS, PARTICIPANT_LABELS)
    except ValueError as exc:
        plan.errors.append(str(exc))
        return plan

    instances = (
        CourseInstance.query.options(joinedload(CourseInstance.course)).all()
    )
    instance_by_id = {i.id: i for i in instances}
    label_to_id = {_participant_event_label(i): i.id for i in instances}

    # Preload проти N+1: реєстрації за reg_id, мапи email->User та
    # (user,instance)->активна реєстрація -- усе для визначення дії
    # (create/update/unchanged) без запитів у циклі.
    #
    # reg_id тягнемо теж: при звичайному циклі "вивантажив -> поправив ->
    # завантажив" кожен рядок має reg_id, і db.session.get у циклі давав
    # запит на кожного учасника.
    reg_ids_in_file = set()
    for r in rows:
        try:
            rid = _int(r.get('reg_id'))
        except ValueError:
            continue  # нечислове значення -- помилка рядка, не префетчу
        if rid:
            reg_ids_in_file.add(rid)
    regs_by_id = {}
    if reg_ids_in_file:
        regs_by_id = {
            r.id: r for r in EventRegistration.query
            .filter(EventRegistration.id.in_(reg_ids_in_file)).all()
        }

    emails_in_file = {
        (_str(r.get('email')) or '').strip().lower()
        for r in rows if _str(r.get('email'))
    }
    users_by_email = {}
    if emails_in_file:
        for u in User.query.filter(User.email.in_(emails_in_file)).all():
            users_by_email[u.email] = u
    active_reg = {}
    if users_by_email:
        uids = [u.id for u in users_by_email.values()]
        for r in (
            EventRegistration.query
            .filter(EventRegistration.user_id.in_(uids),
                    EventRegistration.status != 'cancelled')
            .all()
        ):
            active_reg[(r.user_id, r.instance_id)] = r

    for line_no, raw in enumerate(rows, start=2):
        try:
            reg_id = _int(raw.get('reg_id'))
            reg = regs_by_id.get(reg_id) if reg_id else None
            if reg_id and reg is None:
                raise ValueError(f'реєстрацію id={reg_id} не знайдено')

            last_name = _str(raw.get('last_name'))
            first_name = _str(raw.get('first_name'))
            phone = _str(raw.get('phone'))
            if not last_name:
                raise ValueError('порожнє Прізвище')
            if not first_name:
                raise ValueError("порожнє Ім'я")
            if not phone:
                raise ValueError('порожній Телефон')

            if reg is not None:
                instance_id = reg.instance_id
            else:
                event_label = _str(raw.get('event'))
                instance_id = _resolve_event_to_id(event_label, label_to_id, instance_by_id)
                if instance_id is None:
                    raise ValueError(f'захід {event_label!r} не знайдено')

            ptype_raw = _str(raw.get('participant_type'))
            ptype = PARTICIPANT_TYPE_KEY_BY_LABEL.get(ptype_raw, ptype_raw) if ptype_raw else None
            if ptype and ptype not in VALID_PARTICIPANT_TYPES:
                raise ValueError(f'тип учасника {ptype_raw!r} недопустимий')

            status_raw = _str(raw.get('status')) or 'Підтверджено'
            status = REG_STATUS_KEY_BY_LABEL.get(status_raw, status_raw)
            if status not in VALID_REG_STATUSES:
                raise ValueError(f'статус {status_raw!r} недопустимий')

            pay_raw = _str(raw.get('payment_status')) or 'Не оплачено'
            payment_status = PAYMENT_STATUS_KEY_BY_LABEL.get(pay_raw, pay_raw)
            if payment_status not in VALID_PAYMENT_STATUSES:
                raise ValueError(f'статус оплати {pay_raw!r} недопустимий')

            specs, unknown = _parse_spec_cell(raw.get('specializations'))
            if unknown:
                raise ValueError(f'невідомі спеціалізації: {", ".join(unknown)}')

            # Числові діапазони -- валідуємо тут, щоб дати чітку per-row
            # помилку замість падіння на DB CHECK-constraint у apply.
            payment_amount = _decimal(raw.get('payment_amount'))
            if payment_amount is not None and payment_amount < 0:
                raise ValueError('сума оплати не може бути від\'ємною')
            cpd = _points_cell(raw.get('cpd_points_awarded'))
            if cpd is not None and cpd < 0:
                raise ValueError('бали БПР не можуть бути від\'ємними')
            participation_format = _parse_participation_format(
                raw.get('participation_format'))
            experience = _int(raw.get('experience_years'))
            if experience is not None and not (0 <= experience <= 70):
                raise ValueError('стаж має бути в межах 0-70 років')
            birth = _date(raw.get('birth_date'))
            if birth is not None and (birth > date.today() or birth.year < 1900):
                raise ValueError('некоректна дата народження')

            data = {
                'instance_id': instance_id,
                'last_name': last_name,
                'first_name': first_name,
                'middle_name': _str(raw.get('middle_name')),
                'email': _str(raw.get('email')),
                'phone': phone,
                'participant_type': ptype,
                'birth_date': birth,
                'education': _str(raw.get('education')),
                'workplace': _str(raw.get('workplace')),
                'position': _str(raw.get('position')),
                'specializations': specs,
                'status': status,
                'payment_status': payment_status,
                'payment_amount': payment_amount,
                'attended': _bool(raw.get('attended')),
                'cpd_points_awarded': cpd,
                'participation_format': participation_format,
                'experience_years': experience,
                'license_number': _str(raw.get('license_number')),
                'admin_notes': _str(raw.get('admin_notes')),
            }

            # Дія для preview: reg_id або наявна активна реєстрація за email
            # -> update (або unchanged, якщо нічого не змінюється); інакше create.
            existing_reg = reg
            if existing_reg is None:
                email = (data['email'] or '').strip().lower()
                if email:
                    u = users_by_email.get(email)
                    if u is not None:
                        existing_reg = active_reg.get((u.id, instance_id))

            if existing_reg is not None:
                diff = _diff_participant(existing_reg, data)
                action = 'update' if diff else 'unchanged'
            else:
                diff = []
                action = 'create'

            name = f'{last_name} {first_name}'.strip()
            event_lbl = (
                _participant_event_label(instance_by_id[instance_id])
                if instance_id in instance_by_id else ''
            )
            plan.participants.append({
                'data': data, 'reg_id': reg.id if reg else None, 'action': action,
            })
            plan.changes.append(ParticipantChange(
                action=action, name=name, event=event_lbl, fields_changed=diff,
            ))
        except Exception as exc:
            record_row_error(plan, 'participants', line_no, exc, 'Учасники')
            plan.changes.append(ParticipantChange(
                action='error',
                name=_str(raw.get('last_name')) or f'#{line_no}',
                event=_str(raw.get('event')) or '',
                error=str(exc),
            ))

    return plan


def apply_participants_plan(plan: ParticipantsImportPlan) -> dict:
    """Atomic upsert учасників. Очікує plan.is_valid==True."""
    if not plan.is_valid:
        return {'ok': False, 'reason': 'plan has errors'}

    from app.services import participant_service

    created = 0
    updated = 0
    skipped = 0
    try:
        for item in plan.participants:
            if item.get('action') == 'unchanged':
                skipped += 1
                continue
            reg_id = item['reg_id']
            reg = db.session.get(EventRegistration, reg_id) if reg_id else None
            _reg, was_created = participant_service.upsert_participant(
                item['data'], reg=reg, on_duplicate='update',
            )
            if was_created:
                created += 1
            else:
                updated += 1
        db.session.commit()
        return {'ok': True, 'created': created, 'updated': updated, 'skipped': skipped}
    except participant_service.ParticipantError as exc:
        # Доменна помилка -- її текст написаний для людини, показуємо як є.
        db.session.rollback()
        return {'ok': False, 'reason': str(exc)}
    except Exception:
        db.session.rollback()
        logger.exception('apply_participants_plan failed')
        return {'ok': False, 'reason': _APPLY_FAILED_MESSAGE}


