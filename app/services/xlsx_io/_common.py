"""Спільна основа xlsx-обміну: стилі, ширини, перетворювачі клітинок,
читання аркуша, довідники й тимчасові файли.

Тут немає нічого доменного. Якщо додаєш сюди щось, що знає про курси,
проведення чи учасників -- воно належить сусідньому модулю.
"""
from __future__ import annotations

import io
import logging
import re
import time
import uuid
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from zoneinfo import ZoneInfo
from decimal import Decimal, InvalidOperation
from pathlib import Path

from flask import current_app
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from sqlalchemy.orm import joinedload, selectinload

from app.extensions import db
from app.models.course import Course
from app.models.course_instance import CourseInstance
from app.models.medical_profile import MedicalProfile
from app.models.media_file import MediaFile
from app.models.program_block import ProgramBlock
from app.models.registration import EventRegistration
from app.models.specializations import SPECIALIZATIONS
from app.models.trainer import Trainer
from app.models.user import User
from app.services import trainer_links
from app.utils import ensure_utc

logger = logging.getLogger(__name__)


def _resolve_media_id(url):
    """URL-зображення (/media/...) -> id MediaFile у реєстрі, або None.

    Після Фази 6 курси зберігають зображення лише через media_id. У xlsx
    лишається людиночитний URL; на імпорті резолвимо його назад у реєстр за
    file_path. Невідомі/зовнішні URL -> None (адмін довантажує через UI)."""
    url = (url or '').strip()
    prefix = '/media/'
    if not url.startswith(prefix):
        return None
    file_path = url[len(prefix):]
    media = MediaFile.query.filter_by(file_path=file_path).first()
    return media.id if media else None


# ----------------------------------------------------------------------
# Загальні константи / утиліти
# ----------------------------------------------------------------------

KYIV = ZoneInfo('Europe/Kyiv')

HEADER_FILL = PatternFill('solid', fgColor='4F46E5')
HEADER_FONT = Font(color='FFFFFF', bold=True)
WRAP = Alignment(wrap_text=True, vertical='top')

# ----- Number formats ------------------------------------------------
# Дроблені формати для різних типів даних. Use cell.number_format = ...
FMT_INT = '0'
FMT_POINTS = '0.##'
FMT_CURRENCY_UAH = '#,##0 "₴"'
FMT_DATETIME = 'YYYY-MM-DD HH:MM'
FMT_DATE = 'DD.MM.YYYY'

# Per-key number_format мапа. Якщо ключ відсутній — формат не виставляємо
# (текстовий за замовчуванням).
NUMBER_FORMATS = {
    # Courses
    'id': FMT_INT,
    'base_price': FMT_CURRENCY_UAH,
    'cpd_points_online': FMT_POINTS,
    'cpd_points_offline': FMT_POINTS,
    'max_participants': FMT_INT,
    # Instances
    'price': FMT_CURRENCY_UAH,
    'start_date': FMT_DATETIME,
    'end_date': FMT_DATETIME,
    # Program blocks
    'sort_order': FMT_INT,
    # Participants
    'reg_id': FMT_INT,
    'birth_date': FMT_DATE,
    # Списки адмінки (реєстрації, користувачі)
    'event_date': FMT_DATE,
    'created_at': FMT_DATETIME,
    'last_login_at': FMT_DATETIME,
    'place_number': FMT_INT,
    'user_id': FMT_INT,
    'registrations': FMT_INT,
    'issued_at': FMT_DATETIME,
    'updated_at': FMT_DATETIME,
    'sent_at': FMT_DATETIME,
    'retry_count': FMT_INT,
    'resolved_at': FMT_DATETIME,
    'error_code': FMT_INT,
    'seats_left': FMT_INT,
    'used_count': FMT_INT,
    'max_uses': FMT_INT,
    'per_user_limit': FMT_INT,
    'valid_from': FMT_DATE,
    'valid_until': FMT_DATE,
    'payment_amount': FMT_CURRENCY_UAH,
    'refunded_amount': FMT_CURRENCY_UAH,
    'discount_amount': FMT_CURRENCY_UAH,
    'cpd_points_awarded': FMT_POINTS,
    'experience_years': FMT_INT,
}

# ----- Color fills для enum-полів ------------------------------------


def _fill(hex_color: str) -> PatternFill:
    return PatternFill('solid', fgColor=hex_color)


EVENT_FORMAT_FILLS = {
    'online': _fill('DBEAFE'),       # blue
    'offline': _fill('D1FAE5'),      # green
    'hybrid': _fill('E9D5FF'),       # purple
}

STATUS_FILLS = {
    'draft': _fill('F3F4F6'),        # gray
    'published': _fill('DBEAFE'),    # blue
    'active': _fill('D1FAE5'),       # green
    'completed': _fill('A7F3D0'),    # darker green
    'cancelled': _fill('FECACA'),    # red
}

BOOL_TRUE_FILL = _fill('D1FAE5')     # light green
BOOL_FALSE_FILL = _fill('FEE2E2')    # light red

# Ledь-помітна зебра для непарних data-рядків. Робимо її вручну (а не
# через Excel TableStyle), бо вбудовані стилі дають занадто помітне
# банінг -- ledь-помітної опції серед них нема.
ZEBRA_FILL = _fill('FAFAFA')

# ----- Column widths (ширина = "шт. символів"). ----------------------
# Дають xlsx-у форму "зручний для перегляду", не "ALL DEFAULT 14".
COURSE_WIDTHS = {
    'id': 6,
    'slug': 32,
    'title': 55,
    'subtitle': 40,
    'short_description': 50,
    'description': 60,
    'event_type': 16,
    'base_price': 14,
    'cpd_points_online': 14,
    'cpd_points_offline': 14,
    'max_participants': 12,
    'trainer_slugs': 36,
    'hero_image': 50,
    'card_image': 50,
    'agenda': 40,
    'final_cta_text': 50,
    'target_audience': 50,
    'tags': 28,
    'is_active': 12,
    'is_featured': 14,
}

INSTANCE_WIDTHS = {
    'id': 6,
    'course_slug': 32,
    'topic': 42,
    'start_date': 22,
    'end_date': 22,
    'event_format': 14,
    'price': 14,
    'cpd_points_online': 14,
    'cpd_points_offline': 14,
    'max_participants': 12,
    'trainer_slugs': 36,
    'location': 18,
    'online_link': 40,
    'status': 14,
}

PROGRAM_WIDTHS = {
    'course_slug': 32,
    'sort_order': 10,
    'heading': 40,
    'items': 70,
}

FAQ_WIDTHS = {
    'course_slug': 32,
    'question': 50,
    'answer': 70,
}

TRAINER_WIDTHS = {'slug': 28, 'full_name': 36, 'role': 50}

VALID_FORMATS = {t[0] for t in CourseInstance.FORMATS}
VALID_STATUSES = {t[0] for t in CourseInstance.STATUSES}

FORMAT_LABEL = dict(CourseInstance.FORMATS)  # 'online' -> 'Онлайн' тощо
FORMAT_KEY_BY_LABEL = {v: k for k, v in FORMAT_LABEL.items()}

STATUS_LABEL = dict(CourseInstance.STATUSES)  # 'draft' -> 'Чернетка' тощо
STATUS_KEY_BY_LABEL = {v: k for k, v in STATUS_LABEL.items()}


def _import_dir() -> Path:
    """instance/xlsx_imports -- лежить поза static, недоступне з вебу."""
    inst = Path(current_app.instance_path)
    target = inst / 'xlsx_imports'
    target.mkdir(parents=True, exist_ok=True)
    return target


def save_uploaded_xlsx(file_storage) -> str:
    """Зберегти upload з UI у тимчасову директорію, повернути token."""
    token = uuid.uuid4().hex
    path = _import_dir() / f'{token}.xlsx'
    file_storage.save(str(path))
    return token


def get_uploaded_path(token: str) -> Path | None:
    """Знайти збережений файл за token. Захист від path traversal."""
    if not token.isalnum() or len(token) != 32:
        return None
    p = _import_dir() / f'{token}.xlsx'
    return p if p.is_file() else None


def cleanup_upload(token: str) -> None:
    p = get_uploaded_path(token)
    if p is not None:
        try:
            p.unlink()
        except OSError:
            logger.exception('Failed to remove temp xlsx %s', p)


def cleanup_stale_xlsx_uploads(max_age_minutes: int = 30) -> int:
    """Видалити завантажені для preview xlsx-файли, старші за `max_age_minutes`.

    Викликається з APScheduler-job-у. Повертає кількість видалених файлів
    (для логування).
    """
    target = _import_dir()
    cutoff = datetime.now().timestamp() - max_age_minutes * 60
    removed = 0
    for p in target.glob('*.xlsx'):
        try:
            if p.stat().st_mtime < cutoff:
                p.unlink()
                removed += 1
        except OSError:
            logger.exception('Failed to inspect/remove stale xlsx %s', p)
    if removed:
        logger.info('Cleaned up %d stale xlsx upload(s)', removed)
    return removed


def _style_header(ws, columns: list[str], labels: dict[str, str] | None = None) -> None:
    """Записати заголовки (без виставлення ширин — ширини окремо через
    `_set_column_widths`, бо вони залежать від типу контенту).
    """
    for col_idx, key in enumerate(columns, start=1):
        display = labels.get(key, key) if labels else key
        cell = ws.cell(row=1, column=col_idx, value=display)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical='center', horizontal='left')
    ws.freeze_panes = 'A2'
    # Висота header-рядка для зручності
    ws.row_dimensions[1].height = 22


def _set_column_widths(ws, columns: list[str], widths: dict[str, int]) -> None:
    """Виставити ширину кожної колонки за мапою `widths`. Якщо ключа немає,
    використовуємо дефолт 14."""
    for col_idx, key in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = widths.get(key, 14)


def _apply_number_formats(ws, columns: list[str], last_row: int) -> None:
    """Призначити number_format на кожну колонку (одразу на всі дата-клітинки
    від рядка 2 до last_row), якщо key є в NUMBER_FORMATS."""
    if last_row < 2:
        return
    for col_idx, key in enumerate(columns, start=1):
        fmt = NUMBER_FORMATS.get(key)
        if not fmt:
            continue
        for r in range(2, last_row + 1):
            ws.cell(row=r, column=col_idx).number_format = fmt


def _apply_table_style(ws, columns: list[str], table_name: str, last_data_row: int) -> None:
    """Перетворити діапазон A1:<last_col><last_data_row> на Excel-Table.

    Дає авто-фільтри в заголовку + іменований range. Зебру самі малюємо
    через `_apply_zebra` (бо вбудоване Excel-банінг -- надто помітне).
    """
    if last_data_row < 2:
        return
    last_col = get_column_letter(len(columns))
    ref = f'A1:{last_col}{last_data_row}'
    table = Table(displayName=table_name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name='TableStyleLight1',   # майже-білий, без сильних кольорів
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=False,      # ВЛАСНА зебра нижче
        showColumnStripes=False,
    )
    ws.add_table(table)


def _apply_zebra(ws, n_cols: int, first_data_row: int, last_data_row: int) -> None:
    """Заповнити кожен 2-й data-рядок ледь помітним сірим. Викликати
    ПЕРЕД призначенням enum-fills, щоб кольорові клітинки (event_type,
    status, is_active, ...) перекривали zebra-fill своїм кольором.
    """
    if last_data_row < first_data_row:
        return
    # Колір на другому, четвертому, шостому data-рядку
    # (visually -- стовпчик `1st row=white, 2nd row=gray, 3rd=white...`).
    for row in range(first_data_row + 1, last_data_row + 1, 2):
        for col in range(1, n_cols + 1):
            ws.cell(row, col).fill = ZEBRA_FILL


# Excel не приймає керуючі символи (крім \t \n \r): один \x07 у нотатці
# піднімав IllegalCharacterError і валив ВЕСЬ експорт, а не рядок.
_CONTROL_CHARS_RE = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')
# Ліміт довжини комірки в xlsx; довший рядок Excel вважає файл пошкодженим.
MAX_CELL_LENGTH = 32767


def _safe_text(value: str) -> str:
    """Прибрати керуючі символи й обрізати до ліміту комірки."""
    text = _CONTROL_CHARS_RE.sub('', value)
    if len(text) > MAX_CELL_LENGTH:
        text = text[:MAX_CELL_LENGTH - 1] + '…'
    return text


def write_cell(ws, row, column, value):
    """Записати комірку так, щоб дані лишились ДАНИМИ.

    openpyxl визначає тип за вмістом рядка: значення, що починається з '=',
    стає ЖИВОЮ формулою (data_type 'f'), а '#N/A' та решта ERROR_CODES --
    коміркою-помилкою ('e'). У звіти йде чужий текст (місце роботи, нотатки,
    відповіді SMTP), тобто учасник міг би керувати вмістом файлу, який
    відкриє менеджер. Повертаємо тип у 's': значення зберігається дослівно,
    Excel показує текст.
    """
    if isinstance(value, str):
        value = _safe_text(value)
    cell = ws.cell(row=row, column=column, value=value)
    if cell.data_type in ('f', 'e'):
        cell.data_type = 's'
    return cell


def _to_kyiv_naive(dt):
    """Зняти TZ, попередньо переконвертувавши в Київ. openpyxl попереджує
    про tz-aware datetimes; у клітинці маємо bare datetime, який Excel
    розуміє як «локальний час»."""
    if dt is None:
        return None
    if dt.tzinfo is not None:
        return dt.astimezone(KYIV).replace(tzinfo=None)
    return dt


def _str(v) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _to_lines(value: list | None) -> str:
    """JSON list -> текст по 1 елементу на рядок."""
    if not value:
        return ''
    return '\n'.join(str(x) for x in value)


def _from_lines(value: str | None) -> list[str]:
    if not value:
        return []
    return [line.strip() for line in str(value).splitlines() if line.strip()]


def event_type_dropdown_options():
    """Назви активних типів для drop-down у згенерованому файлі."""
    from app.services import event_types
    return [name for _code, name in event_types.choices()]


def normalize_event_type(raw):
    """Код виду заходу з того, що написали у клітинці.

    Приймає і внутрішній код ('seminar'), і українську назву з drop-down
    ('Семінар'), і застарілий тип: старі вигрузки мусять заходити далі.
    """
    from app.services import event_types

    value = _str(raw) or ''
    if not value:
        return None

    rows = event_types.directory()
    if value in rows:
        return value

    by_label = {row.name: code for code, row in rows.items()}
    if value in by_label:
        return by_label[value]

    # Резервний пошук без урахування регістру: файл редагується руками,
    # і користувач міг набрати назву в іншому регістрі.
    value_cf = value.casefold()
    for code in rows:
        if code.casefold() == value_cf:
            return code
    for name, code in by_label.items():
        if name.casefold() == value_cf:
            return code

    allowed = sorted(rows) + sorted(by_label)
    raise ValueError(f'event_type={value!r} – допустимі: {allowed}')


def _bool(v) -> bool:
    if isinstance(v, bool):
        return v
    if v is None:
        return False
    s = str(v).strip().lower()
    return s in {'true', 'yes', 'y', '1', 'так', 'да'}


def _decimal(v) -> Decimal | None:
    if v is None or v == '':
        return None
    try:
        value = Decimal(str(v).replace(' ', '').replace(',', '.'))
    except (InvalidOperation, ValueError):
        raise ValueError(f'не число: {v!r}')
    # Decimal вважає 'NaN' і 'Infinity' валідними значеннями, і вони проходять
    # далі мовчки: NaN підриває будь-яке порівняння (`NaN < 0` кидає
    # InvalidOperation вже в чужому місці), а Infinity порівняння ПРОХОДИТЬ
    # і доживає до commit, де падає вся транзакція через одну клітинку.
    # `parse_points` в app/utils.py давно робить цю перевірку -- тут її бракувало.
    if not value.is_finite():
        raise ValueError(f'не число: {v!r}')
    return value


def _int(v) -> int | None:
    if v is None or v == '':
        return None
    try:
        # OverflowError, а не ValueError: float('1e400') -> inf, і int(inf)
        # кидає саме його. Без нього нагору йшов англомовний текст
        # "cannot convert float infinity to integer" замість нашого рядка.
        return int(float(str(v).replace(' ', '').replace(',', '.')))
    except (TypeError, ValueError, OverflowError):
        raise ValueError(f'не ціле число: {v!r}')


def _points_cell(v) -> 'Decimal | None':
    """Бали БПР із комірки: приймає і «7,5», і «7.5» (Excel з укр. локаллю)."""
    from app.utils import parse_points

    try:
        return parse_points(v)
    except ValueError:
        raise ValueError(f'некоректні бали БПР: {v!r}')


# Текст, який бачить людина, коли застосування плану впало з несподіваної
# причини. Сирий str(exc) сюди віддавати не можна: у ньому буває фрагмент SQL
# і назви колонок, тобто внутрішня будова БД у відповіді на завантажений файл.
# Повний traceback іде в лог через logger.exception.
_APPLY_FAILED_MESSAGE = (
    'Не вдалося застосувати імпорт. Дані не змінено. '
    'Подробиці записано в лог -- зверніться до адміністратора.'
)


# ----------------------------------------------------------------------
# Межі значень: ОДНЕ місце для того, що дублювалося між розбором xlsx і
# CHECK-обмеженнями моделей.
#
# Досі розбір перевіряв тільки `base_price < 0` на курсі, а решту меж знала
# лише БД. Наслідок: рядок із max_participants=0 або відʼємними балами
# проходив розбір, доходив до commit() і валив УВЕСЬ імпорт -- хоча для будь-
# якої іншої помилки людина отримувала акуратне "Рядок 7: ...". Тепер обидва
# аркуші звіряються з цим переліком, і кожне порушення лишається помилкою
# СВОГО рядка.
#
# Межі звірені з CheckConstraint моделей (app/models/course.py:104-116,
# app/models/course_instance.py:74-87) І з розрядністю самих колонок.
#
# Верхня межа потрібна не менше за нижню, і з менш очевидної причини:
# `Decimal('1e400')` -- цілком СКІНЧЕННЕ число, тож перевірка is_finite() його
# пропускає, порівняння з нулем воно проходить, а Numeric(10,2) у Postgres
# відхиляє з "numeric field overflow" вже на commit -- тобто знову весь імпорт
# гине через одну клітинку.
#
# Numeric(10,2) -> 8 цілих розрядів -> 99 999 999.99
# Numeric(5,2)  -> 3 цілих розряди  -> 999.99
# Integer       -> int4             -> 2 147 483 647
# Додаючи сюди поле, звіряйся з оголошенням колонки, а не з памʼяттю.
_NUMERIC_10_2_MAX = Decimal('99999999.99')
_NUMERIC_5_2_MAX = Decimal('999.99')
_INT4_MAX = 2147483647

_VALUE_LIMITS = {
    'base_price': (Decimal(0), _NUMERIC_10_2_MAX, 'ціна'),
    'price': (Decimal(0), _NUMERIC_10_2_MAX, 'ціна'),
    'cpd_points_online': (Decimal(0), _NUMERIC_5_2_MAX, 'бали БПР (онлайн)'),
    'cpd_points_offline': (Decimal(0), _NUMERIC_5_2_MAX, 'бали БПР (офлайн)'),
    'max_participants': (1, _INT4_MAX, 'місць'),
}


def _check_min_values(parsed: dict) -> None:
    """Кинути ValueError на перше поле, що виходить за межі колонки.

    Порожнє значення (None) пропускаємо: обмеження моделей усі мають форму
    "... OR IS NULL", крім base_price, який розбір і так завжди заповнює.
    """
    for field, (minimum, maximum, label) in _VALUE_LIMITS.items():
        value = parsed.get(field)
        if value is None:
            continue
        if value < minimum:
            raise ValueError(
                f'{label}: значення {value} менше за дозволене {minimum}'
            )
        if value > maximum:
            raise ValueError(
                f'{label}: значення {value} більше за дозволене {maximum}'
            )


def _dt(v) -> datetime | None:
    """Прийняти або datetime (openpyxl auto-parses), або ISO-рядок."""
    if v is None or v == '':
        return None
    if isinstance(v, datetime):
        dt = v
    else:
        try:
            dt = datetime.fromisoformat(str(v))
        except ValueError:
            raise ValueError(f'неможливо розпарсити дату: {v!r}')
    if dt.tzinfo is None:
        # Наївний час у клітинці вважаємо київським. ZoneInfo сам візьме
        # правильний зсув для тієї дати (узимку +2, влітку +3) -- фіксоване
        # число тут було б помилкою двічі на рік.
        dt = dt.replace(tzinfo=KYIV)
    return dt


# Роздільник переліку тренерів у клітинці. Експорт зʼєднує через '; ' (кома
# трапляється ВСЕРЕДИНІ самого ПІБ, напр. «Іванов І. І., PhD», тож нею не
# можна розділяти елементи списку), а на імпорті приймаємо і ';', і ',' --
# людина, що редагує клітинку вручну, радше поставить кому.
#
# Але не одночасно: якщо в клітинці є хоч одна ';', це саме той формат,
# що дає export, і кома в ній -- частина ПІБ, а не роздільник між
# тренерами ("Іванов І. І., PhD; Петров П. П." має розпастись на ДВА
# тренери, а не на три фрагменти). Кома як роздільник -- лише запасний
# варіант для клітинки без жодної ';' узагалі, набраної вручну.
def _split_trainer_names(text: str) -> list[str]:
    sep = ';' if ';' in text else ','
    return [p.strip() for p in text.split(sep) if p.strip()]


def build_trainer_lookup(trainers) -> tuple[dict, dict, set]:
    """Довідники «slug -> id» і «ПІБ -> id» плюс множина неоднозначних ПІБ.

    ПІБ не унікальний у БД, тож словник за іменем мовчки лишав би останнього
    тезку: клітинка з «Іванов І. І.» прив'язала б довільного з двох, і
    менеджер побачив би це лише на сайті. Тезок збираємо окремо і на імпорті
    вимагаємо slug -- він унікальний.
    """
    by_slug = {t.slug: t.id for t in trainers}
    by_name: dict[str, int] = {}
    ambiguous: set[str] = set()
    for t in trainers:
        if t.full_name in by_name:
            ambiguous.add(t.full_name)
        else:
            by_name[t.full_name] = t.id
    for name in ambiguous:
        by_name.pop(name, None)
    return by_slug, by_name, ambiguous


def _resolve_trainer_ids(raw, trainer_id_by_slug: dict, trainer_id_by_name: dict,
                         ambiguous_names: set | None = None) -> list[int]:
    """Клітинка з переліком тренерів (ПІБ і/або slug, через ';'/',') ->
    id тренерів у порядку запису.

    Порядок -- це роль (перший тренер лектор-головний), тож список, а не
    множина. Усі нерозпізнані значення збираємо в ОДНУ помилку рядка: інакше
    менеджер правив би десятиіменну клітинку по одному імені за раунд.

    Дублікати відсіюємо тут-таки, зберігаючи перше входження: `set_trainers`
    робить те саме на записі, і без дзеркальної поведінки тут клітинка з
    повтореним іменем назавжди показувала б «змінено» у прев'ю -- розібраний
    список ніколи не збігся б із дедуплікованим збереженим.
    """
    text = _str(raw)
    if not text:
        return []
    ambiguous_names = ambiguous_names or set()
    ids: list[int] = []
    seen: set[int] = set()
    unknown: list[str] = []
    ambiguous_hit: list[str] = []
    for name in _split_trainer_names(text):
        if name in ambiguous_names:
            ambiguous_hit.append(name)
            continue
        tid = trainer_id_by_slug.get(name) or trainer_id_by_name.get(name)
        if tid is None:
            unknown.append(name)
        elif tid not in seen:
            seen.add(tid)
            ids.append(tid)
    if ambiguous_hit:
        raise ValueError(
            'кілька тренерів мають однакове ПІБ, вкажіть slug замість імені: '
            + ', '.join(repr(n) for n in ambiguous_hit)
        )
    if unknown:
        raise ValueError(
            'тренерів не знайдено (ні за slug, ні за ПІБ): '
            + ', '.join(repr(n) for n in unknown)
        )
    return ids


# ----------------------------------------------------------------------
# Аркуші, довідники й випадні списки. Лежали під банером COURSES, хоча
# ними користуються ВСІ домени -- банер стояв вище за код, і механічний
# поділ повірив саме йому.
# ----------------------------------------------------------------------
# Назви sheet-ів. Експортуємо в українській, парсинг приймає обидва.
SHEET_ALIASES = {
    'courses': ['Курси', 'Courses'],
    'program_blocks': ['Блоки програми', 'Program blocks'],
    'faq': ['FAQ'],
    'instances': ['Розклад', 'Instances'],
    'participants': ['Учасники', 'Participants'],
    'materials': ['Матеріали', 'Materials'],
}


def _find_sheet(wb, key: str):
    """Знайти sheet за будь-яким з прийнятних псевдонімів."""
    for name in SHEET_ALIASES.get(key, []):
        if name in wb.sheetnames:
            return wb[name]
    return None


# Кількість рядків, на які поширюється data-validation drop-down (тип
# заходу, формат, статус тощо). Менеджер може дописувати нові рядки знизу --
# валідація все одно покриватиме. 500 з запасом.
_DROPDOWN_BUFFER_ROWS = 500
_TRAINERS_SHEET_NAME = 'Тренери'


def _add_trainers_sheet(wb) -> int:
    """Додати reference-sheet з активними тренерами. ПІБ йде в колонці A,
    щоб саме воно потрапляло у drop-down тренерів. Slug -- у колонці B
    (для довідки). Повертає номер останнього рядка з даними."""
    ws = wb.create_sheet(_TRAINERS_SHEET_NAME)
    cols = ['full_name', 'slug', 'role']
    _style_header(
        ws,
        cols,
        {'full_name': 'ПІБ (значення)', 'slug': 'Slug', 'role': 'Посада'},
    )

    trainers = (
        Trainer.query.filter_by(is_active=True)
        .order_by(Trainer.full_name)
        .all()
    )
    for row_idx, t in enumerate(trainers, start=2):
        ws.cell(row=row_idx, column=1, value=t.full_name).alignment = WRAP
        ws.cell(row=row_idx, column=2, value=t.slug)
        ws.cell(row=row_idx, column=3, value=t.role or '').alignment = WRAP

    widths = {'full_name': 36, 'slug': 28, 'role': 50}
    _set_column_widths(ws, cols, widths)
    _apply_zebra(ws, len(cols), first_data_row=2, last_data_row=1 + len(trainers))
    _apply_table_style(ws, cols, 'tblTrainers', last_data_row=1 + len(trainers))
    return 1 + len(trainers)


def _add_inline_dropdown(ws, column_key: str, columns: list[str],
                         options: list[str], last_data_row: int,
                         title: str = '', hint: str = '') -> None:
    """Прикріпити drop-down зі статичним списком значень.

    Використовується для невеликих enum-полів (event_type, формат, статус).
    Excel (OOXML) обмежує КОЖЕН із трьох рядків -- formula1, error, prompt --
    255 символами ОКРЕМО; довший список потребує reference-sheet.
    """
    if not options:
        return
    col_letter = get_column_letter(columns.index(column_key) + 1)
    # Inline-list у formula1 має бути обгорнутий лапками й розділений комами.
    formula = '"' + ','.join(options) + '"'
    # error навмисно короткий і БЕЗ переліку значень: повний список -- лише
    # в prompt (hint), бо formula1/error/prompt ділять один ліміт на рядок
    # (255 симв.), а formula1 із 12 активних видів заходу вже займає 221 --
    # запасу лишається ~34 символи, тобто десь два нові коди. Дублювання
    # переліку в error перше впиралося б у межу.
    error_message = 'Оберіть значення зі списку (стрілочка праворуч клітинки).'
    for part_name, text in (
        ('formula1', formula), ('error', error_message), ('prompt', hint),
    ):
        if len(text) > 255:
            # Мовчки віддати файл, який Excel вважає пошкодженим, гірше, ніж
            # віддати його без цієї випадайки.
            logger.warning(
                'Drop-down для %r пропущено: %s -- %s символів (ліміт Excel '
                '255 на кожен з formula1/error/prompt). Для довших списків '
                'потрібен reference-sheet.',
                column_key, part_name, len(text),
            )
            return
    dv = DataValidation(
        type='list',
        formula1=formula,
        allow_blank=False,
        showDropDown=False,  # False у XML = ПОКАЗУВАТИ стрілочку
        errorStyle='stop',
        error=error_message,
        errorTitle='Невалідне значення',
        prompt=hint,
        promptTitle=title,
    )
    final_row = max(last_data_row, 1) + _DROPDOWN_BUFFER_ROWS
    dv.add(f'{col_letter}2:{col_letter}{final_row}')
    ws.add_data_validation(dv)


# Читання аркуша -- спільне для всіх доменів. Лежало під банером COURSES
# разом із рештою помічників аркушів: банер стояв вище за код.
# Підписи колонок, що змінилися після того, як менеджери вже мали на руках
# експорти. Приймаються на імпорті нарівні з чинними -- інакше перейменування
# підпису мовчки лишало б поле незмінним у файлі, який виглядає цілком
# нормальним. Ключ -- старий підпис, значення -- internal key. Живе тут, а не
# поруч із COURSE_LABELS: єдиний споживач -- _read_sheet, спільний для всіх
# аркушів.
LEGACY_HEADERS = {
    # target_audience звузився до допису: перелік спеціальностей на сторінці
    # збирається з довідника, а не з цього поля.
    'Цільова аудиторія': 'target_audience',
}


def _read_sheet(ws, columns: list[str], labels: dict[str, str] | None = None,
                optional: tuple[str, ...] = ()) -> list[dict]:
    """Прочитати sheet у list[dict].

    Заголовки приймаються або як internal key (англ., 'slug'), або як
    українські labels (з `labels`), для зворотньої сумісності зі старими
    xlsx-файлами.

    `optional` -- колонки, відсутність яких НЕ є помилкою: так у формат можна
    додати нове поле, не ламаючи імпорт файлів, експортованих раніше. Ключів
    відсутніх опційних колонок у рядку не буде взагалі (а не None), щоб
    виклик міг відрізнити "у файлі порожньо" від "колонки не було" і не
    занулив наявне значення.

    Заголовок і рядки читаємо одним проходом ітератора, БЕЗ ws.max_row: у
    read-only режимі він дорівнює None, якщо у файлі немає запису
    <dimension> (так зберігають Google Sheets і частина конвертерів), і
    порівняння з числом падало TypeError на цілком нормальному файлі.
    Побічно це виправляє й інше: перевірка колонок більше не пропускається
    для листа без рядків даних -- раніше лист із чужими заголовками тихо
    читався як "порожньо" замість помилки.
    """
    rows_iter = ws.iter_rows(values_only=True)
    header = next(rows_iter, None)
    if header is None:
        return []

    # accepted: будь-яка валідна назва заголовка -> internal key
    accepted: dict[str, str] = {}
    for key in columns:
        accepted[key] = key
        accepted[key.lower()] = key
        if labels and key in labels:
            ua = labels[key]
            accepted[ua] = key
            accepted[ua.lower()] = key
    for legacy, key in LEGACY_HEADERS.items():
        if key in columns:
            accepted.setdefault(legacy, key)
            accepted.setdefault(legacy.lower(), key)

    col_idx: dict[str, int] = {}
    for i, hv in enumerate(header):
        if hv is None:
            continue
        key = accepted.get(str(hv).strip()) or accepted.get(str(hv).strip().lower())
        if key:
            col_idx[key] = i + 1

    missing = [c for c in columns if c not in col_idx and c not in optional]
    if missing:
        pretty = [(labels.get(k, k) if labels else k) for k in missing]
        raise ValueError(
            f'Sheet "{ws.title}": бракує колонок: {", ".join(pretty)}'
        )

    present = [c for c in columns if c in col_idx]
    rows = []
    for row in rows_iter:
        if not any(v is not None and str(v).strip() for v in row):
            continue  # повністю порожній рядок
        # Рядок може бути коротшим за заголовок (обрізані хвостові порожні
        # клітинки) -- беремо None замість IndexError.
        d = {
            name: (row[col_idx[name] - 1] if col_idx[name] <= len(row) else None)
            for name in present
        }
        rows.append(d)
    return rows


# ----------------------------------------------------------------------
# Повернення помилок у сам файл
# ----------------------------------------------------------------------

ERROR_COLUMN_LABEL = 'Помилка'
_ERROR_FILL = _fill('FEE2E2')


def record_row_error(plan, sheet_key: str, line_no: int, message: str,
                     sheet_label: str) -> None:
    """Записати помилку рядка і в текстовий перелік, і в структурований.

    Текстовий (`plan.errors`) читає людина на сторінці прев'ю; структурований
    (`plan.row_errors`) потрібен, щоб повернути ті самі помилки в КОПІЮ
    завантаженого файлу -- правити десять помилок, звіряючись із екраном,
    незручно рівно настільки, наскільки зручно правити їх у самому рядку.
    """
    plan.errors.append(f'Рядок {line_no} ({sheet_label}): {message}')
    plan.row_errors.append((sheet_key, line_no, str(message)))


def annotate_errors_xlsx(path, row_errors) -> 'io.BytesIO | None':
    """Копія завантаженого файлу з колонкою «Помилка» проти винних рядків.

    Повертає None, якщо помилок рядків немає (тоді й повертати нічого) або
    файл уже не читається -- це допоміжна зручність, і вона не має права
    зронити основний сценарій показу помилок.
    """
    if not row_errors:
        return None
    by_sheet: dict[str, dict[int, list[str]]] = {}
    for sheet_key, line_no, message in row_errors:
        by_sheet.setdefault(sheet_key, {}).setdefault(line_no, []).append(message)
    try:
        wb = load_workbook(filename=str(path), read_only=False, data_only=True)
        for sheet_key, rows in by_sheet.items():
            ws = _find_sheet(wb, sheet_key)
            if ws is None:
                continue
            col = ws.max_column + 1
            header = write_cell(ws, 1, col, ERROR_COLUMN_LABEL)
            header.fill = HEADER_FILL
            header.font = HEADER_FONT
            ws.column_dimensions[get_column_letter(col)].width = 60
            for line_no, messages in rows.items():
                cell = write_cell(ws, line_no, col, '; '.join(messages))
                cell.fill = _ERROR_FILL
                cell.alignment = WRAP
        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
        return out
    except Exception:
        logger.exception('annotate_errors_xlsx failed')
        return None
