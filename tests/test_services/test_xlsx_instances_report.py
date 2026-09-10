"""Звіт реєстру проведень (/admin/instances -> xlsx): вид заходу в колонці.

Це ЗВІТ, а не файл для зворотного завантаження: решта його колонок теж
ефективні (effective_title, effective_price, effective_cpd_for), тож і вид
заходу тут ефективний -- власний вид дати, а порожній відкочується на
курсовий. Round-trip-файл розкладу поводиться навпаки й свідомо
(див. test_xlsx_instances.py).
"""
from datetime import datetime, timedelta, timezone
from uuid import uuid4

from openpyxl import load_workbook

from app.extensions import db
from app.models.course import Course
from app.models.course_instance import CourseInstance
from app.services import xlsx_reports


def _course(event_type='seminar'):
    c = Course(title=f'Курс {uuid4().hex[:4]}', slug=f'r-{uuid4().hex[:6]}',
               is_active=True, base_price=0, event_type=event_type)
    db.session.add(c)
    db.session.commit()
    return c


def _instance(course, **kw):
    kw.setdefault('start_date', datetime.now(timezone.utc) + timedelta(days=30))
    inst = CourseInstance(course_id=course.id, status='published',
                          event_format='offline', **kw)
    db.session.add(inst)
    db.session.commit()
    return inst


def _report_row(instances):
    stream = xlsx_reports.export_instances_report_xlsx(
        instances, {i.id: 0 for i in instances})
    ws = load_workbook(stream).active
    header = [c.value for c in ws[1]]
    # Шапка звіту може мати рядок застосованих фільтрів вище -- шукаємо
    # рядок, у якому стоїть підпис колонки ID.
    if xlsx_reports._INST_REPORT_LABELS['id'] not in header:
        for row in ws.iter_rows(min_row=1, max_row=8, values_only=True):
            if row and xlsx_reports._INST_REPORT_LABELS['id'] in row:
                header = list(row)
                break
    data = [r for r in ws.iter_rows(values_only=True)
            if r and r[0] == instances[0].id]
    return dict(zip(header, data[0]))


def test_report_names_the_event_type_of_the_date(client):
    course = _course('seminar')
    inst = _instance(course, event_type='training')
    row = _report_row([inst])
    assert row[xlsx_reports._INST_REPORT_LABELS['event_type']] == 'Тренінг'


def test_report_falls_back_to_the_course_type(client):
    """Звіт друкує ефективний вид: порожній override -- вид курсу.

    На відміну від аркуша «Розклад», де порожньо мусить лишитись порожнім,
    щоб успадкування пережило round-trip.
    """
    course = _course('seminar')
    inst = _instance(course)
    row = _report_row([inst])
    assert row[xlsx_reports._INST_REPORT_LABELS['event_type']] == 'Семінар'


def test_report_leaves_the_cell_empty_when_nobody_has_a_type(client):
    course = _course(None)
    inst = _instance(course)
    row = _report_row([inst])
    assert not row[xlsx_reports._INST_REPORT_LABELS['event_type']]
