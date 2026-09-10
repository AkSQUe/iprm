"""XLSX: колонка trainer_slugs (перелік тренерів) в ОБОХ sheet-ах --
Курси і Розклад.

Порядок у клітинці -- це роль: перший тренер лектор-головний, і саме в
такому порядку його очікує `trainer_links.set_trainers`. Експорт зʼєднує
ПІБ через '; ' (кома трапляється всередині самого ПІБ, напр. «Іванов
І. І., PhD», тож нею не можна розділяти елементи списку); імпорт
приймає і ';', і ',' -- людина, що редагує клітинку вручну, радше
поставить кому. Drop-down у цій колонці прибрано: Excel не вміє
валідувати клітинку з кількома значеннями.
"""
from datetime import datetime, timedelta, timezone
from uuid import uuid4

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from app.extensions import db
from app.models.course import Course
from app.models.course_instance import CourseInstance
from app.models.trainer import Trainer
from app.services import trainer_links, xlsx_io

KYIV = xlsx_io.KYIV


# --- фікстури -----------------------------------------------------------

def _trainer(**kw):
    kw.setdefault('full_name', f'Тренер {uuid4().hex[:6]}')
    t = Trainer(slug=f'tr-{uuid4().hex[:8]}', **kw)
    db.session.add(t)
    db.session.commit()
    return t


def _course(**kw):
    kw.setdefault('title', f'Курс {uuid4().hex[:4]}')
    c = Course(slug=f'x-{uuid4().hex[:6]}', event_type='course',
               is_active=True, base_price=0, **kw)
    db.session.add(c)
    db.session.commit()
    return c


def _instance(course, **kw):
    kw.setdefault('start_date', datetime.now(timezone.utc) + timedelta(days=30))
    kw.setdefault('status', 'published')
    kw.setdefault('event_format', 'offline')
    inst = CourseInstance(course_id=course.id, **kw)
    db.session.add(inst)
    db.session.commit()
    return inst


# --- helpers: Курси -------------------------------------------------------

def _export_courses_sheet():
    return load_workbook(xlsx_io.export_courses_xlsx())['Курси']


def _course_row(ws, slug):
    header = [c.value for c in ws[1]]
    slug_idx = header.index(xlsx_io.COURSE_LABELS['slug'])
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[slug_idx] == slug:
            return dict(zip(header, row))
    return None


def _write_courses_file(tmp_path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Курси'
    cols = xlsx_io.COURSE_COLS
    ws.append([xlsx_io.COURSE_LABELS[c] for c in cols])
    for row in rows:
        ws.append([row.get(c, '') for c in cols])
    path = tmp_path / f'courses-{uuid4().hex[:6]}.xlsx'
    wb.save(path)
    return path


def _base_course_row(course, **overrides):
    row = {
        'id': course.id,
        'slug': course.slug,
        'title': course.title,
        'event_type': 'Курс',
        'base_price': 0,
        'is_active': True,
        'is_featured': False,
    }
    row.update(overrides)
    return row


# --- helpers: Розклад -----------------------------------------------------

def _export_instances_sheet():
    return load_workbook(xlsx_io.export_instances_xlsx())['Розклад']


def _instance_row(ws, course_slug):
    header = [c.value for c in ws[1]]
    slug_idx = header.index(xlsx_io.INSTANCE_LABELS['course_slug'])
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[slug_idx] == course_slug:
            return dict(zip(header, row))
    return None


def _write_instances_file(tmp_path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Розклад'
    cols = xlsx_io.INSTANCE_COLS
    ws.append([xlsx_io.INSTANCE_LABELS[c] for c in cols])
    for row in rows:
        ws.append([row.get(c, '') for c in cols])
    path = tmp_path / f'sch-{uuid4().hex[:6]}.xlsx'
    wb.save(path)
    return path


def _base_instance_row(course, **overrides):
    row = {
        'course_slug': course.slug,
        'start_date': (datetime.now(timezone.utc) + timedelta(days=30))
        .astimezone(KYIV).replace(tzinfo=None).isoformat(),
        'event_format': 'Офлайн',
        'status': 'Опубліковано',
        'location': 'Харків',
    }
    row.update(overrides)
    return row


# ======================================================================
# КУРСИ
# ======================================================================

def test_course_round_trip_keeps_order_of_several_trainers(client, tmp_path):
    """Порядок з xlsx -> id тренерів у ТОМУ Ж порядку після реекспорту.

    Тренери навмисно заведені в БД у порядку id, що НЕ збігається з
    бажаним порядком викладання (c, a, b, а не a, b, c за зростанням id)
    -- інакше тест пройшов би і при прихованому сортуванні за id, нічого
    насправді не перевіривши.
    """
    course = _course()
    a = _trainer(full_name='Тренер Авдієнко')
    b = _trainer(full_name='Тренер Бондар')
    c = _trainer(full_name='Тренер Величко')
    assert a.id < b.id < c.id  # передумова: id зростають у порядку a,b,c

    trainer_links.set_trainers(course, [c.id, a.id, b.id])
    db.session.commit()

    ws = _export_courses_sheet()
    cell = _course_row(ws, course.slug)[xlsx_io.COURSE_LABELS['trainer_slugs']]
    assert cell == f'{c.full_name}; {a.full_name}; {b.full_name}'

    path = _write_courses_file(tmp_path, [_base_course_row(course, trainer_slugs=cell)])
    plan = xlsx_io.parse_courses_xlsx(path)
    assert plan.is_valid, plan.errors
    parsed = plan.courses[0]['parsed']
    assert parsed['trainer_ids'] == [c.id, a.id, b.id]

    assert xlsx_io.apply_courses_plan(plan)['ok']
    db.session.expire(course)
    assert [t.id for t in course.trainers] == [c.id, a.id, b.id]


def test_course_export_uses_semicolon_separator(client):
    """Кома не годиться -- «Іванов І. І., PhD» цілком можливе ПІБ, і кома
    у ролі роздільника розбила б його навпіл."""
    course = _course()
    t1 = _trainer(full_name='Іванов І. І., PhD')
    t2 = _trainer(full_name='Петров П. П.')
    trainer_links.set_trainers(course, [t1.id, t2.id])
    db.session.commit()

    ws = _export_courses_sheet()
    cell = _course_row(ws, course.slug)[xlsx_io.COURSE_LABELS['trainer_slugs']]
    assert cell == f'{t1.full_name}; {t2.full_name}'
    assert cell.count(';') == 1  # рівно один роздільник між ДВОМА тренерами


def test_course_import_accepts_comma_separator(client, tmp_path):
    """Людина, яка друкує вручну, поставить кому."""
    course = _course()
    t1 = _trainer()
    t2 = _trainer()
    path = _write_courses_file(tmp_path, [
        _base_course_row(course, trainer_slugs=f'{t1.slug}, {t2.slug}'),
    ])
    plan = xlsx_io.parse_courses_xlsx(path)
    assert plan.is_valid, plan.errors
    assert plan.courses[0]['parsed']['trainer_ids'] == [t1.id, t2.id]


def test_course_unknown_trainer_reports_every_bad_value_at_once(client, tmp_path):
    """Одна помилка рядка з переліком усіх нерозпізнаних, а не падіння
    на першому."""
    course = _course()
    path = _write_courses_file(tmp_path, [
        _base_course_row(course, trainer_slugs='Невідомий Перший; Невідомий Другий'),
    ])
    plan = xlsx_io.parse_courses_xlsx(path)
    assert not plan.is_valid
    row_errors = [e for e in plan.errors if 'Невідомий Перший' in e]
    assert len(row_errors) == 1
    assert 'Невідомий Другий' in row_errors[0]


def test_course_import_with_empty_trainer_cell_still_imports(client, tmp_path):
    course = _course()
    path = _write_courses_file(tmp_path, [_base_course_row(course, trainer_slugs='')])
    plan = xlsx_io.parse_courses_xlsx(path)
    assert plan.is_valid, plan.errors
    assert plan.courses[0]['parsed']['trainer_ids'] == []
    assert xlsx_io.apply_courses_plan(plan)['ok']


def test_course_trainer_column_has_no_dropdown(client):
    """Excel не вміє валідувати клітинку з кількома значеннями -- залишений
    drop-down мовчки відхиляв би коректний перелік тренерів."""
    ws = _export_courses_sheet()
    col_idx = xlsx_io.COURSE_COLS.index('trainer_slugs') + 1
    col_letter = get_column_letter(col_idx)
    for dv in ws.data_validations.dataValidation:
        for rng in dv.sqref.ranges:
            assert rng.min_col != col_idx and rng.max_col != col_idx, (
                f'Знайдено data-validation у колонці тренерів ({col_letter})'
            )
    # Reference-sheet з тренерами лишається -- джерело точних написань.
    assert xlsx_io._TRAINERS_SHEET_NAME in ws.parent.sheetnames


# ======================================================================
# РОЗКЛАД
# ======================================================================

def test_instance_round_trip_keeps_order_of_several_trainers(client, tmp_path):
    course = _course()
    inst = _instance(course)
    a = _trainer(full_name='Тренер Авдієнко')
    b = _trainer(full_name='Тренер Бондар')
    c = _trainer(full_name='Тренер Величко')
    assert a.id < b.id < c.id

    trainer_links.set_trainers(inst, [c.id, a.id, b.id])
    db.session.commit()

    ws = _export_instances_sheet()
    cell = _instance_row(ws, course.slug)[xlsx_io.INSTANCE_LABELS['trainer_slugs']]
    assert cell == f'{c.full_name}; {a.full_name}; {b.full_name}'

    path = _write_instances_file(tmp_path, [
        _base_instance_row(course, id=inst.id, trainer_slugs=cell),
    ])
    plan = xlsx_io.parse_instances_xlsx(path)
    assert plan.is_valid, plan.errors
    parsed = plan.instances[0]['parsed']
    assert parsed['trainer_ids'] == [c.id, a.id, b.id]

    assert xlsx_io.apply_instances_plan(plan)['ok']
    db.session.expire(inst)
    assert [t.id for t in inst.trainers] == [c.id, a.id, b.id]


def test_instance_export_uses_semicolon_separator(client):
    course = _course()
    inst = _instance(course)
    t1 = _trainer(full_name='Іванов І. І., PhD')
    t2 = _trainer(full_name='Петров П. П.')
    trainer_links.set_trainers(inst, [t1.id, t2.id])
    db.session.commit()

    ws = _export_instances_sheet()
    cell = _instance_row(ws, course.slug)[xlsx_io.INSTANCE_LABELS['trainer_slugs']]
    assert cell == f'{t1.full_name}; {t2.full_name}'
    assert cell.count(';') == 1


def test_instance_import_accepts_comma_separator(client, tmp_path):
    course = _course()
    t1 = _trainer()
    t2 = _trainer()
    path = _write_instances_file(tmp_path, [
        _base_instance_row(course, trainer_slugs=f'{t1.slug}, {t2.slug}'),
    ])
    plan = xlsx_io.parse_instances_xlsx(path)
    assert plan.is_valid, plan.errors
    assert plan.instances[0]['parsed']['trainer_ids'] == [t1.id, t2.id]


def test_instance_unknown_trainer_reports_every_bad_value_at_once(client, tmp_path):
    course = _course()
    path = _write_instances_file(tmp_path, [
        _base_instance_row(course, trainer_slugs='Невідомий Перший; Невідомий Другий'),
    ])
    plan = xlsx_io.parse_instances_xlsx(path)
    assert not plan.is_valid
    row_errors = [e for e in plan.errors if 'Невідомий Перший' in e]
    assert len(row_errors) == 1
    assert 'Невідомий Другий' in row_errors[0]


def test_instance_import_with_empty_trainer_cell_still_imports(client, tmp_path):
    course = _course()
    path = _write_instances_file(tmp_path, [_base_instance_row(course, trainer_slugs='')])
    plan = xlsx_io.parse_instances_xlsx(path)
    assert plan.is_valid, plan.errors
    assert plan.instances[0]['parsed']['trainer_ids'] == []
    assert xlsx_io.apply_instances_plan(plan)['ok']


def test_instance_trainer_column_has_no_dropdown(client):
    ws = _export_instances_sheet()
    col_idx = xlsx_io.INSTANCE_COLS.index('trainer_slugs') + 1
    for dv in ws.data_validations.dataValidation:
        for rng in dv.sqref.ranges:
            assert rng.min_col != col_idx and rng.max_col != col_idx
    assert xlsx_io._TRAINERS_SHEET_NAME in ws.parent.sheetnames


# ======================================================================
# ЗМІНА ПОРЯДКУ -- ЦЕ ЗМІНА (diff)
# ======================================================================

def test_course_reordering_trainers_is_reported_as_change(client, tmp_path):
    course = _course()
    t1 = _trainer()
    t2 = _trainer()
    trainer_links.set_trainers(course, [t1.id, t2.id])
    db.session.commit()

    path = _write_courses_file(tmp_path, [
        _base_course_row(course, trainer_slugs=f'{t2.slug}, {t1.slug}'),
    ])
    plan = xlsx_io.parse_courses_xlsx(path)
    assert plan.is_valid, plan.errors
    change = next(ch for ch in plan.changes if ch.slug == course.slug)
    assert change.action == 'update'
    assert 'trainer_slugs' in change.fields_changed


def test_instance_reordering_trainers_is_reported_as_change(client, tmp_path):
    course = _course()
    inst = _instance(course)
    t1 = _trainer()
    t2 = _trainer()
    trainer_links.set_trainers(inst, [t1.id, t2.id])
    db.session.commit()

    path = _write_instances_file(tmp_path, [
        _base_instance_row(course, id=inst.id, trainer_slugs=f'{t2.slug}, {t1.slug}'),
    ])
    plan = xlsx_io.parse_instances_xlsx(path)
    assert plan.is_valid, plan.errors
    change = next(ch for ch in plan.changes if ch.course_slug == course.slug)
    assert change.action == 'update'
    assert 'trainer_slugs' in change.fields_changed
