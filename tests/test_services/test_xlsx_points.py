"""Бали БПР у xlsx: комірка приймає укр. кому (Excel з укр. локаллю),
дробовий числовий формат, розщеплені колонки online/offline і формат
участі на імпорті реєстрацій.
"""
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from uuid import uuid4

from openpyxl import Workbook, load_workbook

from app.extensions import db
from app.models.course import Course
from app.models.course_instance import CourseInstance
from app.models.instance_tariff import InstanceTariff
from app.models.registration import EventRegistration
from app.models.user import User
from app.services import xlsx_io
from app.services.xlsx_io import _points_cell


def test_points_cell_accepts_comma():
    assert _points_cell('7,5') == Decimal('7.50')


def test_points_cell_accepts_dot_and_blank():
    assert _points_cell('7.5') == Decimal('7.50')
    assert _points_cell('') is None


def test_points_cell_rejects_text():
    import pytest
    with pytest.raises(ValueError):
        _points_cell('багато')


# --- курси: розщеплені колонки балів ----------------------------------------

def _course(**kw):
    kw.setdefault('title', f'Курс {uuid4().hex[:4]}')
    c = Course(slug=f'pts-{uuid4().hex[:6]}', event_type='course',
               is_active=True, base_price=0, **kw)
    db.session.add(c)
    db.session.commit()
    return c


def test_course_export_has_online_offline_points_columns(client):
    _course(cpd_points_online=Decimal('7.50'), cpd_points_offline=Decimal('5.00'))
    ws = load_workbook(xlsx_io.export_courses_xlsx())['Курси']
    header = [c.value for c in ws[1]]
    assert xlsx_io.COURSE_LABELS['cpd_points_online'] in header
    assert xlsx_io.COURSE_LABELS['cpd_points_offline'] in header
    rows = [dict(zip(header, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
    row = rows[0]
    assert row[xlsx_io.COURSE_LABELS['cpd_points_online']] == 7.5
    assert row[xlsx_io.COURSE_LABELS['cpd_points_offline']] == 5.0


def test_course_import_reads_comma_points(client, tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Курси'
    ws.append([xlsx_io.COURSE_LABELS[c] for c in xlsx_io.COURSE_COLS])
    row = {
        'slug': f'pts-{uuid4().hex[:6]}',
        'title': 'Новий курс',
        'event_type': 'Курс',
        'base_price': '0',
        'cpd_points_online': '7,5',
        'cpd_points_offline': '5,0',
        'is_active': 'Так',
    }
    ws.append([row.get(c, '') for c in xlsx_io.COURSE_COLS])
    path = tmp_path / 'courses.xlsx'
    wb.save(path)

    plan = xlsx_io.parse_courses_xlsx(path)
    assert plan.is_valid, plan.errors
    assert xlsx_io.apply_courses_plan(plan)['ok']

    course = Course.query.filter_by(slug=row['slug']).one()
    assert course.cpd_points_online == Decimal('7.50')
    assert course.cpd_points_offline == Decimal('5.00')


# --- реєстрації: формат участі на імпорті -----------------------------------

def _instance():
    course = Course(title=f'Курс {uuid4().hex[:4]}', slug=f'ptsi-{uuid4().hex[:6]}',
                    is_active=True, event_type='course')
    db.session.add(course)
    db.session.flush()
    inst = CourseInstance(
        course_id=course.id, status='published', event_format='hybrid',
        start_date=datetime.now(timezone.utc) + timedelta(days=30),
    )
    db.session.add(inst)
    db.session.commit()
    return inst


def _write_participants(tmp_path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Учасники'
    ws.append([xlsx_io.PARTICIPANT_LABELS[c] for c in xlsx_io.PARTICIPANT_COLS])
    for row in rows:
        ws.append([row.get(c, '') for c in xlsx_io.PARTICIPANT_COLS])
    path = tmp_path / f'part-{uuid4().hex[:6]}.xlsx'
    wb.save(path)
    return path


def _participant_row(inst, **overrides):
    row = {
        'event': f'#{inst.id}',
        'last_name': 'Тестовий',
        'first_name': 'Тест',
        'email': f'u-{uuid4().hex[:6]}@test.com',
        'phone': '+380501112233',
        'workplace': 'Клініка',
        'status': 'Підтверджено',
        'payment_status': 'Не оплачено',
        'attended': 'Ні',
    }
    row.update(overrides)
    return row


def test_participant_export_shows_effective_participation_format(client):
    inst = _instance()
    user = User.create_with_password(f'e-{uuid4().hex[:6]}@test.com', 'password123',
                                     first_name='Ол', last_name='Ів')
    db.session.flush()
    reg = EventRegistration(
        user_id=user.id, instance_id=inst.id, phone='+380501112233',
        specialty='Дерматолог', workplace='Клініка',
        participation_format='online',
    )
    db.session.add(reg)
    db.session.commit()

    ws = load_workbook(xlsx_io.export_participants_xlsx())['Учасники']
    header = [c.value for c in ws[1]]
    row = dict(zip(header, [c.value for c in ws[2]]))
    assert row[xlsx_io.PARTICIPANT_LABELS['participation_format']] == 'Онлайн'


def test_participant_import_reads_online_label(client, tmp_path):
    inst = _instance()
    path = _write_participants(
        tmp_path, [_participant_row(inst, participation_format='Онлайн')])

    plan = xlsx_io.parse_participants_xlsx(path)
    assert plan.is_valid, plan.errors
    assert xlsx_io.apply_participants_plan(plan)['ok']

    reg = EventRegistration.query.filter_by(instance_id=inst.id).one()
    assert reg.participation_format == 'online'


def test_participant_import_blank_participation_format_is_none(client, tmp_path):
    inst = _instance()
    path = _write_participants(
        tmp_path, [_participant_row(inst, participation_format='')])

    plan = xlsx_io.parse_participants_xlsx(path)
    assert plan.is_valid, plan.errors
    assert xlsx_io.apply_participants_plan(plan)['ok']

    reg = EventRegistration.query.filter_by(instance_id=inst.id).one()
    assert reg.participation_format is None


def test_participant_import_garbage_participation_format_is_row_error(client, tmp_path):
    inst = _instance()
    path = _write_participants(
        tmp_path, [_participant_row(inst, participation_format='Гібрид-сяк-так')])

    plan = xlsx_io.parse_participants_xlsx(path)
    assert not plan.is_valid
    assert any('формат участі' in e.lower() for e in plan.errors)


def test_participant_export_shows_raw_column_not_tariff_derived(client):
    # Формат участі не проставлений (власна колонка NULL), а бали
    # обчислюються з тарифу онлайн: effective_participation_format віддає
    # 'online', але експорт мусить показати "за тарифом" (порожньо),
    # інакше круговорот вивантаження/довантаження зафіксував би похідне
    # значення у власній колонці.
    inst = _instance()
    tariff = InstanceTariff(instance_id=inst.id, name='Онлайн', price=0,
                             event_format='online')
    db.session.add(tariff)
    db.session.flush()
    user = User.create_with_password(f'ex-{uuid4().hex[:6]}@test.com', 'password123',
                                     first_name='Ол', last_name='Ів')
    db.session.flush()
    reg = EventRegistration(
        user_id=user.id, instance_id=inst.id, phone='+380501112233',
        specialty='Дерматолог', workplace='Клініка', tariff_id=tariff.id,
    )
    db.session.add(reg)
    db.session.commit()

    assert reg.participation_format is None
    assert reg.effective_participation_format == 'online'

    ws = load_workbook(xlsx_io.export_participants_xlsx())['Учасники']
    header = [c.value for c in ws[1]]
    row = dict(zip(header, [c.value for c in ws[2]]))
    # Порожня комірка після реального save/load openpyxl повертається як
    # None, а не '' -- саме так, як після справжнього циклу
    # вивантаження/довантаження в Excel.
    assert row[xlsx_io.PARTICIPANT_LABELS['participation_format']] in (None, '')


def test_participant_roundtrip_does_not_freeze_derived_format(client, tmp_path):
    # Круговорот "вивантажив -> нічого не міняв -> завантажив" не повинен
    # записати обчислений з тарифу формат у власну колонку реєстрації --
    # інакше подальша зміна тарифу вже не рухала б бали.
    inst = _instance()
    tariff = InstanceTariff(instance_id=inst.id, name='Онлайн', price=0,
                             event_format='online')
    db.session.add(tariff)
    db.session.flush()
    user = User.create_with_password(f'rt-{uuid4().hex[:6]}@test.com', 'password123',
                                     first_name='Р', last_name='Т')
    db.session.flush()
    reg = EventRegistration(
        user_id=user.id, instance_id=inst.id, phone='+380501112233',
        specialty='Дерматолог', workplace='Клініка', tariff_id=tariff.id,
    )
    db.session.add(reg)
    db.session.commit()
    reg_id = reg.id

    buf = xlsx_io.export_participants_xlsx()
    path = tmp_path / 'roundtrip.xlsx'
    path.write_bytes(buf.getvalue())

    plan = xlsx_io.parse_participants_xlsx(path)
    assert plan.is_valid, plan.errors
    result = xlsx_io.apply_participants_plan(plan)
    assert result['ok']

    reg = db.session.get(EventRegistration, reg_id)
    assert reg.participation_format is None
    assert reg.effective_participation_format == 'online'
