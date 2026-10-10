"""Сторінка "Виручка й зобов'язання" і її xlsx.

Правила звіту -- tests/test_services/test_finance_report.py. Тут: доступ
(фінансовий результат не відкривається кожному, хто бачить реєстрації),
вибір місяця, і що файл несе ті самі суми, що й сторінка.
"""
import io
from datetime import datetime, timedelta, timezone
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from app.extensions import db
from app.models.course import Course
from app.models.course_instance import CourseInstance
from app.models.registration import EventRegistration
from app.models.user import User
from app.rbac.registry import ROLES_BY_NAME
from tests.support.rbac import grant_role, make_super_admin, make_user_with_role


def _uid():
    return uuid4().hex[:8]


def _login(client, user):
    with client.session_transaction() as s:
        s['_user_id'] = str(user.id)


@pytest.fixture
def admin(app):
    return make_super_admin(f'rev-{_uid()}@test.com')


@pytest.fixture
def paid_past_and_future(app):
    """Минулого місяця: захід відбувся (1000) і оплата за захід через 40 днів (1500)."""
    now = datetime.now(timezone.utc)
    last_month = (now.replace(day=1) - timedelta(days=15))
    rows = []
    for ends_at, amount in ((last_month, 1000), (now + timedelta(days=40), 1500)):
        user = User.create_with_password(f'rev-{_uid()}@test.com', 'password123',
                                         first_name='Іван', last_name='Петренко')
        course = Course(title=f'Захід {_uid()}', slug=f'rev-{_uid()}', is_active=True)
        db.session.add(course)
        db.session.flush()
        inst = CourseInstance(course_id=course.id, status='published',
                              event_format='offline', start_date=ends_at, end_date=ends_at)
        db.session.add(inst)
        db.session.flush()
        reg = EventRegistration(
            user_id=user.id, instance_id=inst.id, phone='+380501112233',
            specialty='S', workplace='W', status='confirmed', payment_status='paid',
            payment_method='invoice', payment_amount=amount,
        )
        reg.set_paid_at(last_month - timedelta(days=3), 'datetime')
        db.session.add(reg)
        rows.append(reg)
    db.session.commit()
    return last_month.strftime('%Y-%m'), rows


def test_page_renders_for_super_admin(client, admin, paid_past_and_future):
    month, (done, owed) = paid_past_and_future
    _login(client, admin)

    response = client.get(f'/admin/revenue?month={month}')

    html = response.get_data(as_text=True)
    assert response.status_code == 200
    assert done.instance.effective_title in html
    assert owed.instance.effective_title in html


def test_bad_month_falls_back_to_current(client, admin):
    _login(client, admin)
    assert client.get('/admin/revenue?month=garbage').status_code == 200


def test_manager_without_permission_is_refused(client):
    manager = make_user_with_role('manager', f'rev-{_uid()}@test.com')
    _login(client, manager)
    assert client.get('/admin/revenue').status_code in (302, 403)


def test_viewer_does_not_get_it_by_default():
    """Спостерігач бачить розділ "Продажі", але не фінансовий результат."""
    assert 'revenue.view' not in ROLES_BY_NAME['viewer'].defaults


def test_export_has_both_sheets_with_page_sums(client, admin, paid_past_and_future):
    month, (done, owed) = paid_past_and_future
    _login(client, admin)

    response = client.get(f'/admin/revenue/export?month={month}')

    assert response.status_code == 200
    wb = load_workbook(io.BytesIO(response.data))
    assert {'Виручка', "Зобов'язання", 'Фільтри'} <= set(wb.sheetnames)

    def _orders(sheet):
        ws = wb[sheet]
        header = [c.value for c in ws[1]]
        col = header.index('Замовлення')
        return {row[col] for row in ws.iter_rows(min_row=2, values_only=True)}

    assert f'REG-{done.id}' in _orders('Виручка')
    assert f'REG-{owed.id}' in _orders("Зобов'язання")
    assert f'REG-{owed.id}' not in _orders('Виручка')


def test_export_needs_its_own_permission(client):
    user = make_user_with_role('viewer', f'rev-{_uid()}@test.com')
    _login(client, user)
    assert client.get('/admin/revenue/export').status_code in (302, 403)


def test_withdrawal_waiting_for_decision_is_named(client, admin, paid_past_and_future):
    """Відмова, чия заявка на повернення чекає рішення, -- зобов'язання, і
    сторінка каже чому, а не «дату не задано»."""
    from app.models.refund_request import RefundRequest

    month, (done, _owed) = paid_past_and_future
    done.status = 'cancelled'
    db.session.add(RefundRequest(registration_id=done.id, user_id=done.user_id,
                                 reason='Не зможу', status='new', quoted_code='standard'))
    db.session.commit()
    _login(client, admin)

    html = client.get(f'/admin/revenue?month={month}').get_data(as_text=True)

    assert 'відмова учасника' in html
    assert 'чекає рішення про повернення' in html
