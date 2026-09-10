"""Поле speaker_info прибрано з коду: його заміщає блок спікерів із карток.

На публічній сторінці speaker_info ніколи не рендерилось -- блок спікерів
завжди збирався з карток тренерів, тож ручний допис лише дублював і
розходився з карткою. Ці тести фіксують відсутність поля в усіх точках, де
воно раніше читалось/писалось (адмінська форма, партнерське API, xlsx-
експорт), і що старий xlsx зі стовпцем speaker_info і далі імпортується --
менеджер не має лишитись зі зламаним файлом просто тому, що колонку прибрали.
"""
from uuid import uuid4

import pytest
from openpyxl import Workbook, load_workbook

from tests.support.rbac import grant_role

from app.extensions import db
from app.models.course import Course
from app.models.course_instance import CourseInstance
from app.models.site_settings import SiteSettings
from app.models.user import User
from app.services import xlsx_io

API_KEY = 'speaker-info-removed-test-key'


def _uid():
    return uuid4().hex[:8]


@pytest.fixture
def admin_client(app, client):
    user = User.create_with_password(
        f'spkinfo-admin-{_uid()}@test.com', 'Passw0rd!123',
        first_name='Адмін', email_confirmed=True,
    )
    user.is_active = True
    grant_role(user, 'super_admin')
    db.session.commit()
    with client.session_transaction() as session:
        session['_user_id'] = str(user.id)
    yield client
    # Прибираємо за собою: база живе через увесь прогін тестів.
    db.session.delete(user)
    db.session.commit()


@pytest.fixture
def course(app):
    c = Course(title=f'Курс {_uid()}', slug=f'spkinfo-{_uid()}',
               event_type='course', base_price=0, is_active=True)
    db.session.add(c)
    db.session.commit()
    yield c
    db.session.rollback()
    db.session.delete(db.session.merge(c))
    db.session.commit()


@pytest.fixture
def partner_enabled(app):
    settings = SiteSettings.get()
    settings.partner_integration_enabled = True
    settings.partner_api_key = API_KEY
    db.session.commit()
    yield settings
    settings.partner_integration_enabled = False
    settings.partner_api_key = ''
    db.session.commit()


def test_course_form_has_no_speaker_info_field(admin_client, course):
    resp = admin_client.get(f'/admin/courses/{course.id}/edit')
    assert resp.status_code == 200
    body = resp.get_data(as_text=True)
    # Сторінка справді відрендерилась з формою курсу -- інакше відсутність
    # 'speaker_info' нижче нічого б не доводила (могла впасти на 500).
    assert 'name="agenda"' in body
    assert 'speaker_info' not in body
    assert 'Інформація про спікера' not in body


def test_api_event_detail_has_no_speaker_info_key(client, partner_enabled, course):
    inst = CourseInstance(
        course_id=course.id, status='published', event_format='offline',
        price=0,
    )
    db.session.add(inst)
    db.session.commit()

    resp = client.get(f'/api/v1/events/{course.slug}',
                      headers={'X-API-Key': API_KEY})
    assert resp.status_code == 200
    data = resp.get_json()
    # Відповідь справді про потрібний курс і має інші поля detail --
    # інакше відсутність ключа нижче могла б означати лише порожню відповідь.
    assert data['slug'] == course.slug
    assert 'agenda' in data
    assert 'speaker_info' not in data


def test_xlsx_export_has_no_speaker_info_column(client, course):
    wb = load_workbook(xlsx_io.export_courses_xlsx())
    ws = wb['Курси']
    header = [c.value for c in ws[1]]
    # Заголовок справді прочитався (не порожній лист) -- перевіряємо
    # присутність сусідньої, свідомо збереженої колонки.
    assert 'Програма (опис)' in header
    assert 'speaker_info' not in header
    assert 'Інфо про спікера' not in header


def test_xlsx_import_tolerates_files_with_the_old_column(client, tmp_path, course):
    """Зайві колонки імпорт ігнорує -- старі файли лишаються робочими."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Курси'
    cols = list(xlsx_io.COURSE_COLS) + ['speaker_info']
    labels = dict(xlsx_io.COURSE_LABELS, speaker_info='Інфо про спікера')
    ws.append([labels[c] for c in cols])
    row = {
        'id': course.id, 'slug': course.slug, 'title': 'Оновлена назва',
        'event_type': 'Курс', 'base_price': 0, 'is_active': True,
        'is_featured': False, 'speaker_info': 'Старий текст про спікера',
    }
    ws.append([row.get(c, '') for c in cols])
    path = tmp_path / f'courses-old-{_uid()}.xlsx'
    wb.save(path)

    plan = xlsx_io.parse_courses_xlsx(path)
    assert plan.is_valid, plan.errors
    result = xlsx_io.apply_courses_plan(plan)
    assert result['ok']
    assert db.session.get(Course, course.id).title == 'Оновлена назва'
